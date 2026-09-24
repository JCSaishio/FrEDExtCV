"""
FrED Fiber Measure (v7 - diameter native to the laptop)
=======================================================

Real-time fiber diameter measurement from a USB camera, and remote control of
automated FrED experiments over WiFi.

Since v7 the diameter never leaves this laptop while it is measured:

* Every camera frame is grabbed on its own thread and timestamped the moment
  it arrives, then measured on a second thread. The GUI only displays results
  (one feed at a time), so the measurement rate is the camera's own frame
  rate instead of being held back by drawing the video.
* The diameter is graphed live here, under the video. FrED's Pi graphs only
  the temperature and the spooler speed.
* During a FrED experiment the Pi records temperature/spooler data on its
  clock and this app records the diameter on the laptop clock. Retrieve Data
  merges both onto one timeline into the same CSV + Excel files as before.

Time synchronisation (diameter vs. temperature / spooler)
---------------------------------------------------------
1. Each machine timestamps its own samples with its own monotonic clock at the
   moment of acquisition: the laptop when a camera frame arrives
   (``time.perf_counter``), the Pi when its control tick reads the sensors.
2. While connected, the laptop pings the Pi every 2 s (a few bytes). From the
   four timestamps of each ping/reply (NTP method) it measures the offset
   between the two clocks with an error bound of half the round trip. Only
   the fastest round trip of every 10 s is kept (network queuing only ever
   adds delay), and a straight line fitted through those also removes the
   slow drift between the two clock crystals.
3. t = 0 is the instant FrED starts RECORDING (by its timer or the START
   RECORDING NOW button). The Pi reports that instant on its own clock; the
   laptop maps it onto the laptop clock with the fit above, so both data sets
   start at the same instant regardless of WiFi delays. The estimated error
   (typically a few ms) is written to the "Run info" sheet.
4. The two streams keep their own rates: FrED writes one row per control tick
   (e.g. 50 Hz); the camera delivers ~30 fps. Each FrED row gets the latest
   camera frame at or before its time - never interpolated. "Diameter new
   frame" is 1 when that frame appears for the first time and 0 when it is a
   repeat that only fills the row, and "Diameter camera frame #" names it.
   Every frame, at the full camera rate, is kept in its own sheet / CSV.

Run with:  python fiber_measure.py
"""

import base64
import bisect
import csv
import datetime as _dt
import io
import json
import math
import os
import queue
import socket
import statistics
import threading
import time
import tkinter as tk
from collections import deque
from tkinter import filedialog, font as tkfont, messagebox, simpledialog, ttk

import cv2
import numpy as np
from PIL import Image, ImageTk

APP_VERSION = "v7"

# --------------------------------------------------------------------------- #
# Paths
# --------------------------------------------------------------------------- #
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.path.join(BASE_DIR, "Data")
CALIB_FILE = os.path.join(BASE_DIR, "calibration.json")

# --------------------------------------------------------------------------- #
# FrED Pi hotspot defaults (must match laptop_link.py / setup_hotspot.sh)
# --------------------------------------------------------------------------- #
PI_HOTSPOT_SSID = "FrED_Pi"
PI_HOTSPOT_PASSWORD = "fredfiber123"
PI_DEFAULT_IP = "192.168.4.1"   # the Pi's address while acting as the hotspot
PI_DEFAULT_PORT = 5005          # TCP port the Pi listens on

# Spooling-motor PID gain limits (Kp, Ki, Kd). Must match
# UserInterface.MOTOR_GAIN_LIMITS on the Pi, which also clamps to them.
MOTOR_GAIN_LIMITS = (1.0, 15.0, 0.05)
SAMPLE_RATE_LIMITS = (1.0, 100.0)   # FrED control + data rate (Hz)

# Seconds between a frame being exposed and OpenCV handing it over (USB
# transfer + driver buffering, typically ~1 frame). Frames are stamped on
# arrival; if you measure your camera's delay (e.g. film an LED switched at a
# known time), put it here to shift the camera timeline earlier by that much.
CAMERA_LATENCY_S = 0.0

CAMERA_RESOLUTIONS = {
    "Camera default": None,
    "640 x 480": (640, 480),
    "1280 x 720": (1280, 720),
    "1920 x 1080": (1920, 1080),
}

# Diameter jitter filter (same strength as the Pi's former filter): a median
# over FILTER_MEDIAN detected frames removes single-frame spikes, then a mean
# over FILTER_MEAN medians smooths the rest. The live graph uses trailing
# windows; the exported data uses CENTRED windows, which add no time lag, so
# the filtered diameter stays aligned with FrED's data.
FILTER_MEDIAN = 5
FILTER_MEAN = 3

SYNC_INTERVAL_S = 2.0   # clock-sync ping period while connected
STALE_FRAME_S = 0.5     # a FrED row gets no diameter if the newest camera
                        # frame is older than this (camera stopped)

# Real-world units the calibration may use, in millimetres.
UNIT_TO_MM = {"mm": 1.0, "um": 0.001, "µm": 0.001, "micron": 0.001,
              "microns": 0.001, "cm": 10.0, "m": 1000.0, "in": 25.4}

# Column order shared by the laptop-only recording CSV and Excel exports.
EXPORT_FIELDS = [
    "timestamp", "elapsed_s", "frame", "diameter_px", "diameter_real",
    "units", "min_px", "max_px", "std_px", "length_px", "angle_deg",
]

# Columns written as whole numbers in the experiment export.
INT_COLUMNS = {"Temp new reading", "Spooler new reading", "Diameter new frame",
               "Diameter camera frame #", "Camera frame #", "Fiber detected",
               "Steady state"}


# --------------------------------------------------------------------------- #
# Calibration
# --------------------------------------------------------------------------- #
class Calibration:
    """Stores and persists the pixel -> real-world conversion factor."""

    def __init__(self):
        # units_per_pixel: how many real-world units one pixel represents.
        self.units_per_pixel = None  # None -> not calibrated (report in px)
        self.units = "mm"
        self.load()

    @property
    def is_calibrated(self):
        return self.units_per_pixel is not None and self.units_per_pixel > 0

    def to_real(self, value_px):
        """Convert a pixel measurement to real units (or return px if not set)."""
        if self.is_calibrated:
            return value_px * self.units_per_pixel
        return value_px

    def display_scale(self):
        """(factor, unit) that turns pixels into the graphed/exported unit:
        millimetres whenever the calibration unit converts to mm, else the
        calibration's own unit, else pixels."""
        if not self.is_calibrated:
            return 1.0, "px"
        to_mm = UNIT_TO_MM.get(str(self.units).strip().lower())
        if to_mm is None:
            return self.units_per_pixel, self.units
        return self.units_per_pixel * to_mm, "mm"

    def set_from_reference(self, measured_px, known_real, units):
        if measured_px <= 0:
            raise ValueError("Measured pixel size must be > 0.")
        self.units_per_pixel = float(known_real) / float(measured_px)
        self.units = units
        self.save()

    def set_factor(self, units_per_pixel, units):
        self.units_per_pixel = float(units_per_pixel)
        self.units = units
        self.save()

    def clear(self):
        self.units_per_pixel = None
        self.save()

    def load(self):
        try:
            with open(CALIB_FILE, "r", encoding="utf-8") as fh:
                data = json.load(fh)
            self.units_per_pixel = data.get("units_per_pixel")
            self.units = data.get("units", "mm")
        except (FileNotFoundError, json.JSONDecodeError, ValueError):
            pass

    def save(self):
        data = {"units_per_pixel": self.units_per_pixel, "units": self.units}
        try:
            with open(CALIB_FILE, "w", encoding="utf-8") as fh:
                json.dump(data, fh, indent=2)
        except OSError:
            pass


# --------------------------------------------------------------------------- #
# WiFi (TCP) command link to the FrED Raspberry Pi
# --------------------------------------------------------------------------- #
class FredLink:
    """WiFi command link to the FrED Pi. The diameter is NOT sent.

    The Pi (``laptop_link.py``) is the server; this is the client. Messages are
    newline-delimited JSON objects with a ``type`` field (experiment,
    start_now, abort, get_data, sync). Replies are parsed on a background
    thread and queued for the Tk main loop; clock-sync replies carry the laptop
    time they arrived (``_t4``).
    """

    CONNECT_TIMEOUT = 5.0   # seconds to wait when opening the connection

    def __init__(self):
        self.sock = None
        self.host = None
        self.port = None
        self.rx_queue = queue.Queue()
        self._reader = None
        self._stop = threading.Event()
        self._send_lock = threading.Lock()

    @property
    def is_open(self):
        return self.sock is not None

    def open(self, host, port):
        """Connect to the Pi at host:port; raises on failure."""
        self.close()
        sock = socket.create_connection((host, int(port)),
                                        timeout=self.CONNECT_TIMEOUT)
        sock.setsockopt(socket.IPPROTO_TCP, socket.TCP_NODELAY, 1)
        self.sock = sock
        self.host = host
        self.port = int(port)
        # Start the background reader for Pi -> laptop messages.
        self._stop.clear()
        self._reader = threading.Thread(target=self._reader_loop, daemon=True)
        self._reader.start()

    def close(self):
        self._stop.set()
        if self.sock is not None:
            try:
                self.sock.close()
            except Exception:
                pass
        self.sock = None

    @property
    def endpoint(self):
        if self.host is None:
            return "?"
        return f"{self.host}:{self.port}"

    def _reader_loop(self):
        """Read newline-delimited JSON from the Pi and queue each message."""
        buffer = bytearray()
        sock = self.sock
        if sock is None:
            return
        sock.settimeout(0.5)
        while not self._stop.is_set():
            try:
                data = sock.recv(65536)
            except socket.timeout:
                continue
            except Exception:
                break
            if not data:        # Pi closed the connection
                break
            t_arrival = time.perf_counter()
            search_from = len(buffer)
            buffer += data
            newline = buffer.find(b"\n", search_from)
            while newline != -1:
                line = bytes(buffer[:newline])
                del buffer[:newline + 1]
                self._queue_line(line, t_arrival)
                newline = buffer.find(b"\n")
        # The Pi closed the connection (or a read error): mark the link
        # closed so is_open reflects reality. Only close if the socket is
        # still OUR socket (a reconnect may already have installed a new one).
        if not self._stop.is_set() and self.sock is sock:
            self.close()

    def _queue_line(self, line, t_arrival):
        text = line.decode("utf-8", errors="ignore").strip()
        if not text:
            return
        try:
            msg = json.loads(text)
        except (ValueError, json.JSONDecodeError):
            return
        if not isinstance(msg, dict):
            return
        if msg.get("type") == "sync_reply":
            msg["_t4"] = t_arrival
        self.rx_queue.put(msg)

    def send(self, obj):
        """Send one JSON object; closes the link on error. Returns success."""
        if not self.is_open:
            return False
        try:
            with self._send_lock:
                self.sock.sendall((json.dumps(obj) + "\n").encode("utf-8"))
            return True
        except Exception:
            self.close()
            return False

    def ping(self, ping_id):
        """Send a clock-sync ping stamped with the laptop clock."""
        return self.send({"type": "sync", "id": ping_id,
                          "t1": time.perf_counter()})


# --------------------------------------------------------------------------- #
# Clock synchronisation (laptop clock <-> FrED Pi clock)
# --------------------------------------------------------------------------- #
class ClockSync:
    """Estimate the FrED Pi clock from ping/reply timestamps (NTP method).

    For each exchange: t1 ping sent (laptop), t2 ping received (Pi), t3 reply
    sent (Pi), t4 reply received (laptop).
        offset = ((t2 - t1) + (t3 - t4)) / 2      (Pi clock - laptop clock)
        delay  = (t4 - t1) - (t3 - t2)            (network round trip)
    The offset error is at most delay / 2, and queuing only ever ADDS delay,
    so the fastest exchange of every BIN_S seconds is kept. A line
    offset = a + b * t through those points removes the slow drift between the
    two clock crystals (b is the drift, typically a few ppm).
    """

    BIN_S = 10.0
    FIT_MIN_SPAN_S = 30.0
    MAX_SAMPLES = 50000

    def __init__(self):
        self._lock = threading.Lock()
        self._samples = []        # (boot, t_mid_laptop, offset, delay)
        self.last_boot = None

    def add_reply(self, msg):
        try:
            t1 = float(msg["t1"])
            t2 = float(msg["t2"])
            t3 = float(msg["t3"])
            t4 = float(msg["_t4"])
        except (KeyError, TypeError, ValueError):
            return
        delay = (t4 - t1) - (t3 - t2)
        if delay < 0 or delay > 5.0:
            return
        offset = ((t2 - t1) + (t3 - t4)) / 2.0
        boot = msg.get("boot")
        with self._lock:
            self._samples.append((boot, (t1 + t4) / 2.0, offset, delay))
            if len(self._samples) > self.MAX_SAMPLES:
                del self._samples[:len(self._samples) - self.MAX_SAMPLES]
        self.last_boot = boot

    def estimate(self, boot=None):
        """Best offset model for the Pi session ``boot`` (default: latest).

        Returns a dict (a, b, err, n, points, delay_min, boot) or None.
        """
        boot = boot if boot is not None else self.last_boot
        with self._lock:
            samples = [s for s in self._samples if s[0] == boot]
        if not samples:
            return None
        best_per_bin = {}
        for _, t_mid, offset, delay in samples:
            key = int(t_mid // self.BIN_S)
            if key not in best_per_bin or delay < best_per_bin[key][2]:
                best_per_bin[key] = (t_mid, offset, delay)
        points = sorted(best_per_bin.values())
        delay_min = min(p[2] for p in points)
        good = [p for p in points if p[2] <= 3.0 * delay_min + 0.005]
        t = np.array([p[0] for p in good])
        off = np.array([p[1] for p in good])
        if len(good) >= 3 and t[-1] - t[0] >= self.FIT_MIN_SPAN_S:
            t_ref = t[0]
            b, a_ref = np.polyfit(t - t_ref, off, 1)
            a = a_ref - b * t_ref
            resid = off - (a_ref + b * (t - t_ref))
            rms = float(np.sqrt(np.mean(resid ** 2)))
        else:
            fastest = sorted(good, key=lambda p: p[2])[:3]
            a = float(np.median([p[1] for p in fastest]))
            b = 0.0
            rms = 0.0
        return {"a": float(a), "b": float(b), "err": delay_min / 2.0 + rms,
                "n": len(samples), "points": len(good),
                "delay_min": delay_min, "boot": boot}

    @staticmethod
    def to_laptop(est, t_pi):
        """Map a Pi-clock time onto the laptop clock with estimate ``est``."""
        # t_pi = t_laptop + (a + b * t_laptop)
        return (t_pi - est["a"]) / (1.0 + est["b"])


# --------------------------------------------------------------------------- #
# Fiber detection / measurement
# --------------------------------------------------------------------------- #
class FiberDetector:
    """
    Detects a bright, elongated fiber on a dark background and measures its
    diameter (the short dimension), tolerating tilt.

    Pipeline:
        grayscale -> Gaussian blur -> threshold (Otsu or manual) ->
        morphological clean-up -> largest contour -> minAreaRect for
        orientation -> rotate the fiber's own strip flat -> per-column
        thickness profile.

    Speed (v7): only the fiber's minAreaRect strip is rotated (a few thousand
    pixels instead of the whole frame), and the annotated/mask images are
    drawn only for frames that are actually displayed.
    """

    # Preset detection parameters (used at start-up and by "Reset to defaults").
    DEFAULTS = {
        "use_otsu": True,       # auto threshold
        "manual_thresh": 110,   # used when use_otsu is False
        "blur_ksize": 5,        # odd kernel size for Gaussian blur
        "min_area": 500,        # ignore blobs smaller than this (px^2)
    }
    KERNEL = np.ones((3, 3), np.uint8)
    STRIP_MARGIN = 6            # px added around the fiber strip when rotating

    def __init__(self):
        self.use_otsu = FiberDetector.DEFAULTS["use_otsu"]
        self.manual_thresh = FiberDetector.DEFAULTS["manual_thresh"]
        self.blur_ksize = FiberDetector.DEFAULTS["blur_ksize"]
        self.min_area = FiberDetector.DEFAULTS["min_area"]

    def _threshold(self, gray):
        k = max(1, self.blur_ksize)
        if k % 2 == 0:
            k += 1
        blur = cv2.GaussianBlur(gray, (k, k), 0)
        if self.use_otsu:
            _, mask = cv2.threshold(
                blur, 0, 255, cv2.THRESH_BINARY + cv2.THRESH_OTSU
            )
        else:
            _, mask = cv2.threshold(
                blur, self.manual_thresh, 255, cv2.THRESH_BINARY
            )
        # Clean up: remove specks, then fill small gaps inside the fiber.
        mask = cv2.morphologyEx(mask, cv2.MORPH_OPEN, self.KERNEL, iterations=1)
        mask = cv2.morphologyEx(mask, cv2.MORPH_CLOSE, self.KERNEL, iterations=2)
        return mask

    def process(self, frame, want_views=True):
        """
        Returns (annotated_bgr_frame, mask_view, result_dict).

        ``mask_view`` is the binary detection image (after blur/threshold/
        morphology) with the detected contour and diameter drawn on it, so the
        processed feed shows exactly where edge detection happens and how the
        parameters affect it. Both images are None when ``want_views`` is
        False (frames that are measured but not displayed).

        result_dict keys:
            found (bool), diameter_px, min_px, max_px, std_px, length_px,
            angle_deg, n_samples
        """
        gray = cv2.cvtColor(frame, cv2.COLOR_BGR2GRAY)
        mask = self._threshold(gray)

        contours, _ = cv2.findContours(
            mask, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE
        )
        annotated = mask_view = None
        if want_views:
            annotated = frame.copy()
            # Colour copy of the mask so we can draw coloured overlays on it.
            mask_view = cv2.cvtColor(mask, cv2.COLOR_GRAY2BGR)
        result = {
            "found": False,
            "diameter_px": 0.0,
            "min_px": 0.0,
            "max_px": 0.0,
            "std_px": 0.0,
            "length_px": 0.0,
            "angle_deg": 0.0,
            "n_samples": 0,
        }

        if not contours:
            return annotated, mask_view, result

        largest = max(contours, key=cv2.contourArea)
        if cv2.contourArea(largest) < self.min_area:
            # Show the rejected blob outline in red so the Min-area effect is visible.
            if want_views:
                cv2.drawContours(mask_view, [largest], 0, (0, 0, 255), 2)
            return annotated, mask_view, result

        rect = cv2.minAreaRect(largest)  # ((cx,cy),(w,h),angle)
        (cx, cy), (rw, rh), angle = rect

        # Orient so the longer side (fiber length) becomes horizontal.
        if rw < rh:
            rot_angle = angle + 90.0
            length_guess = rh
        else:
            rot_angle = angle
            length_guess = rw

        profile = self._thickness_profile(mask, (cx, cy), rot_angle,
                                          length_guess, min(rw, rh))

        if profile.size == 0:
            return annotated, mask_view, result

        diameter = float(np.median(profile))
        result.update(
            found=True,
            diameter_px=diameter,
            min_px=float(np.min(profile)),
            max_px=float(np.max(profile)),
            std_px=float(np.std(profile)),
            length_px=float(length_guess),
            angle_deg=float(rot_angle),
            n_samples=int(profile.size),
        )

        if not want_views:
            return annotated, mask_view, result

        # ---- overlay (drawn on both the live frame and the mask view) ---- #
        box = cv2.boxPoints(rect).astype(np.int32)
        theta = np.deg2rad(rot_angle)
        px, py = -np.sin(theta), np.cos(theta)  # perpendicular direction
        half = diameter / 2.0
        p1 = (int(cx - px * half), int(cy - py * half))
        p2 = (int(cx + px * half), int(cy + py * half))
        for canvas in (annotated, mask_view):
            cv2.drawContours(canvas, [box], 0, (0, 255, 0), 2)
            # Short perpendicular tick at the fiber centre showing the diameter.
            cv2.line(canvas, p1, p2, (0, 0, 255), 2)
            cv2.circle(canvas, (int(cx), int(cy)), 3, (0, 0, 255), -1)

        return annotated, mask_view, result

    @classmethod
    def _thickness_profile(cls, mask, center, rot_angle, length, thickness):
        """
        Rotate the fiber's strip so the fiber is horizontal, then count white
        pixels in each column to build a thickness (diameter) profile in px.

        Only the minAreaRect strip (plus a margin) is rendered: warpAffine
        computes just the destination pixels, so the cost scales with the
        fiber's own area instead of the whole frame's.
        """
        cx, cy = center
        out_w = int(math.ceil(length)) + 2 * cls.STRIP_MARGIN
        out_h = int(math.ceil(thickness)) + 2 * cls.STRIP_MARGIN
        M = cv2.getRotationMatrix2D(center, rot_angle, 1.0)
        # Rotate about the fiber centre, then move that centre to the middle
        # of the output strip. The shift is a whole number of pixels so the
        # strip samples exactly the pixels a full-frame rotation would.
        M[0, 2] += round(out_w / 2.0 - cx)
        M[1, 2] += round(out_h / 2.0 - cy)
        rotated = cv2.warpAffine(
            mask, M, (out_w, out_h), flags=cv2.INTER_NEAREST, borderValue=0
        )
        col_counts = np.count_nonzero(rotated, axis=0).astype(np.float32)

        nz = np.flatnonzero(col_counts > 0)
        if nz.size == 0:
            return np.array([], dtype=np.float32)

        # Trim the outer 10% of the fiber span to avoid tapered ends.
        x0, x1 = nz[0], nz[-1]
        span = x1 - x0
        if span > 20:
            margin = int(span * 0.10)
            x0 += margin
            x1 -= margin
        profile = col_counts[x0 : x1 + 1]
        profile = profile[profile > 0]
        return profile


# --------------------------------------------------------------------------- #
# Camera capture + measurement threads
# --------------------------------------------------------------------------- #
class CameraWorker:
    """Grab and measure EVERY camera frame on two background threads.

    * capture thread: grabs frames as fast as the camera delivers them and
      stamps each one with the laptop clock the moment it arrives;
    * measure thread: runs the detector on every frame, in order, and hands
      the result to ``on_measurement(frame_no, t, result)``.

    The GUI asks for a display image with :meth:`request_view` and collects
    it with :meth:`take_view`; overlays are only drawn for those frames.
    If measuring ever falls behind the camera, the oldest queued frames are
    dropped (counted in ``dropped``) rather than letting the delay grow.
    """

    QUEUE_FRAMES = 8

    def __init__(self, detector, on_measurement):
        self.detector = detector
        self.on_measurement = on_measurement
        self.cap = None
        self._frames = deque()
        self._cond = threading.Condition()
        self._stop = threading.Event()
        self._threads = []
        self._view_lock = threading.Lock()
        self._view_kind = None     # "live" / "mask" wanted for the next frame
        self._view = None          # rendered view waiting for the GUI
        self.running = False
        self.error = None
        self.frame_no = 0
        self.dropped = 0
        self.capture_fps = 0.0
        self.measure_fps = 0.0
        self.width = self.height = 0
        self.fourcc = ""

    def start(self, index, resolution=None, high_fps=False):
        """(Re)open camera ``index``; returns True if it is delivering."""
        self.stop()
        # CAP_DSHOW avoids slow start-up on Windows.
        cap = cv2.VideoCapture(index, cv2.CAP_DSHOW)
        if not cap.isOpened():
            cap.release()
            cap = cv2.VideoCapture(index)
        if not cap.isOpened():
            cap.release()
            return False
        if high_fps:
            cap.set(cv2.CAP_PROP_FOURCC, cv2.VideoWriter_fourcc(*"MJPG"))
        if resolution:
            cap.set(cv2.CAP_PROP_FRAME_WIDTH, resolution[0])
            cap.set(cv2.CAP_PROP_FRAME_HEIGHT, resolution[1])
        if high_fps:
            cap.set(cv2.CAP_PROP_FPS, 120)   # the driver picks its fastest
        self.cap = cap
        self.width = int(cap.get(cv2.CAP_PROP_FRAME_WIDTH))
        self.height = int(cap.get(cv2.CAP_PROP_FRAME_HEIGHT))
        code = int(cap.get(cv2.CAP_PROP_FOURCC))
        self.fourcc = "".join(chr((code >> 8 * i) & 0xFF) for i in range(4))
        if not self.fourcc.isprintable():
            self.fourcc = ""
        self.error = None
        self.dropped = 0
        self.capture_fps = self.measure_fps = 0.0
        self._frames.clear()
        self._stop.clear()
        self._threads = [
            threading.Thread(target=self._capture_loop, daemon=True),
            threading.Thread(target=self._measure_loop, daemon=True),
        ]
        for th in self._threads:
            th.start()
        self.running = True
        return True

    def stop(self):
        self._stop.set()
        with self._cond:
            self._cond.notify_all()
        for th in self._threads:
            th.join(timeout=2.0)
        self._threads = []
        if self.cap is not None:
            self.cap.release()
            self.cap = None
        self.running = False

    def request_view(self, kind):
        """Ask for the next measured frame to be rendered as ``kind``."""
        with self._view_lock:
            if self._view is None and self._view_kind is None:
                self._view_kind = kind

    def take_view(self):
        """Return the latest rendered view (BGR) once, or None."""
        with self._view_lock:
            view, self._view = self._view, None
        return view

    def _capture_loop(self):
        cap = self.cap
        failures = 0
        count, t_rate = 0, time.perf_counter()
        while not self._stop.is_set():
            if not cap.grab():
                failures += 1
                if failures > 100:
                    self.error = "The camera stopped delivering frames."
                    break
                time.sleep(0.01)
                continue
            t_frame = time.perf_counter() - CAMERA_LATENCY_S
            ok, frame = cap.retrieve()
            if not ok or frame is None:
                continue
            failures = 0
            self.frame_no += 1
            with self._cond:
                if len(self._frames) >= self.QUEUE_FRAMES:
                    self._frames.popleft()
                    self.dropped += 1
                self._frames.append((self.frame_no, t_frame, frame))
                self._cond.notify()
            count += 1
            now = time.perf_counter()
            if now - t_rate >= 1.0:
                self.capture_fps = count / (now - t_rate)
                count, t_rate = 0, now

    def _measure_loop(self):
        count, t_rate = 0, time.perf_counter()
        while not self._stop.is_set():
            with self._cond:
                while not self._frames and not self._stop.is_set():
                    self._cond.wait(0.2)
                if self._stop.is_set():
                    break
                frame_no, t_frame, frame = self._frames.popleft()
            with self._view_lock:
                kind = self._view_kind
            try:
                annotated, mask_view, result = self.detector.process(
                    frame, want_views=kind is not None)
            except Exception as exc:
                print(f"[CameraWorker] detection error: {exc}")
                continue
            if kind is not None:
                with self._view_lock:
                    self._view = annotated if kind == "live" else mask_view
                    self._view_kind = None
            try:
                self.on_measurement(frame_no, t_frame, result)
            except Exception as exc:
                print(f"[CameraWorker] measurement callback error: {exc}")
            count += 1
            now = time.perf_counter()
            if now - t_rate >= 1.0:
                self.measure_fps = count / (now - t_rate)
                count, t_rate = 0, now


# --------------------------------------------------------------------------- #
# Measurement log (every frame, laptop clock)
# --------------------------------------------------------------------------- #
class MeasurementLog:
    """Every measured frame on the laptop clock (thread-safe, append-only).

    One tuple per frame:
        (t, frame_no, found, d_px, min_px, max_px, std_px, factor, unit,
         filtered_live)
    ``factor`` * px = value in ``unit`` (mm when calibrated in a metric unit).
    """

    T, FRAME, FOUND, D_PX, MIN_PX, MAX_PX, STD_PX, FACTOR, UNIT, FILT = range(10)

    def __init__(self):
        self._lock = threading.Lock()
        self._rows = []
        self._times = []

    def append(self, row):
        with self._lock:
            self._rows.append(row)
            self._times.append(row[0])

    def since(self, t_start, t_end=None):
        with self._lock:
            i = bisect.bisect_left(self._times, t_start)
            j = (len(self._times) if t_end is None
                 else bisect.bisect_right(self._times, t_end))
            return self._rows[i:j]

    def last(self):
        with self._lock:
            return self._rows[-1] if self._rows else None

    def first_time(self):
        with self._lock:
            return self._times[0] if self._times else None

    def trim_before(self, t_keep):
        with self._lock:
            i = bisect.bisect_left(self._times, t_keep)
            if i > 0:
                del self._rows[:i]
                del self._times[:i]


def centered_filter(values, median_n=FILTER_MEDIAN, mean_n=FILTER_MEAN):
    """Zero-lag jitter filter: centred median, then centred mean."""
    v = np.asarray(values, dtype=float)
    if v.size == 0:
        return v
    from numpy.lib.stride_tricks import sliding_window_view
    half = median_n // 2
    med = np.median(sliding_window_view(np.pad(v, half, mode="edge"),
                                        median_n), axis=1)
    half = mean_n // 2
    return sliding_window_view(np.pad(med, half, mode="edge"),
                               mean_n).mean(axis=1)


# --------------------------------------------------------------------------- #
# Live graph (plain Tk canvas - cheap enough to never slow the camera)
# --------------------------------------------------------------------------- #
class LivePlot:
    """Small live line graph drawn on a Tk canvas (no matplotlib)."""

    PAD_L, PAD_R, PAD_T, PAD_B = 64, 16, 12, 36
    GRID = "#e6e6e6"
    AXIS = "#444444"

    def __init__(self, parent):
        self.canvas = tk.Canvas(parent, background="white",
                                highlightthickness=1,
                                highlightbackground="#c8c8c8")

    @staticmethod
    def _ticks(lo, hi, target=6):
        span = hi - lo
        if span <= 0:
            return [lo], 1.0
        raw = span / target
        mag = 10 ** math.floor(math.log10(raw))
        step = mag
        for m in (1, 2, 2.5, 5, 10):
            step = m * mag
            if span / step <= target:
                break
        first = math.ceil(lo / step) * step
        ticks = []
        value = first
        while value <= hi + step * 1e-6:
            ticks.append(value)
            value += step
        return ticks, step

    @staticmethod
    def _decimals(step):
        """Decimals needed to label ticks spaced ``step`` apart."""
        if step >= 1:
            return 0
        exponent = math.floor(math.log10(step))
        mantissa = step / 10 ** exponent
        return min(6, -exponent + (1 if abs(mantissa - 2.5) < 1e-6 else 0))

    def draw(self, x_min, x_max, series, x_label, y_label,
             hlines=(), vlines=(), empty_text="Waiting for the fiber..."):
        """Redraw everything.

        series: [(xs, ys, color, width, label)]
        hlines: [(y, color, label)]   (dashed horizontal lines, e.g. target)
        vlines: [(x, color, label)]   (event markers, e.g. recording start)
        """
        c = self.canvas
        c.delete("all")
        w, h = c.winfo_width(), c.winfo_height()
        if w < 80 or h < 60:
            return
        x0, y0 = self.PAD_L, self.PAD_T
        x1, y1 = w - self.PAD_R, h - self.PAD_B
        if x_max <= x_min:
            x_max = x_min + 1.0

        ys_all = [y for s in series for y in s[1]]
        if not ys_all:
            c.create_rectangle(x0, y0, x1, y1, outline=self.AXIS)
            c.create_text((x0 + x1) / 2, (y0 + y1) / 2, text=empty_text,
                          fill="#888888", font=("Segoe UI", 12))
            return
        lo, hi = min(ys_all), max(ys_all)
        for y, _, _ in hlines:
            # Keep a target line in view only if it is near the data.
            if lo - (hi - lo + 1e-9) * 2 <= y <= hi + (hi - lo + 1e-9) * 2:
                lo, hi = min(lo, y), max(hi, y)
        span = hi - lo
        if span < 1e-9:
            span = max(abs(hi) * 0.02, 1e-3)
        lo -= span * 0.08
        hi += span * 0.08

        def sx(x):
            return x0 + (x - x_min) / (x_max - x_min) * (x1 - x0)

        def sy(y):
            return y1 - (y - lo) / (hi - lo) * (y1 - y0)

        # Grid + tick labels
        yt, ystep = self._ticks(lo, hi, 6)
        yd = self._decimals(ystep)
        for y in yt:
            py = sy(y)
            c.create_line(x0, py, x1, py, fill=self.GRID)
            c.create_text(x0 - 6, py, text=f"{y:.{yd}f}", anchor="e",
                          fill=self.AXIS, font=("Segoe UI", 9))
        xt, xstep = self._ticks(x_min, x_max, 8)
        xd = self._decimals(xstep)
        for x in xt:
            px = sx(x)
            c.create_line(px, y0, px, y1, fill=self.GRID)
            c.create_text(px, y1 + 4, text=f"{x:.{xd}f}", anchor="n",
                          fill=self.AXIS, font=("Segoe UI", 9))
        c.create_rectangle(x0, y0, x1, y1, outline=self.AXIS)
        c.create_text((x0 + x1) / 2, h - 4, text=x_label, anchor="s",
                      fill=self.AXIS, font=("Segoe UI", 9))
        c.create_text(12, (y0 + y1) / 2, text=y_label, angle=90,
                      fill=self.AXIS, font=("Segoe UI", 9))

        for y, color, label in hlines:
            if lo <= y <= hi:
                py = sy(y)
                c.create_line(x0, py, x1, py, fill=color, dash=(6, 4), width=2)
                c.create_text(x1 - 4, py - 3, text=label, anchor="se",
                              fill=color, font=("Segoe UI", 9))

        for xs, ys, color, width, _ in series:
            if len(xs) < 2:
                continue
            coords = []
            for x, y in zip(xs, ys):
                coords.append(sx(x))
                coords.append(sy(y))
            c.create_line(*coords, fill=color, width=width)

        for x, color, label in vlines:
            if x_min <= x <= x_max:
                px = sx(x)
                c.create_line(px, y0, px, y1, fill=color, width=2)
                c.create_text(px + 4, y0 + 2, text=label, anchor="nw",
                              fill=color, font=("Segoe UI", 9, "bold"))

        # Legend
        lx = x0 + 8
        for _, _, color, width, label in series:
            c.create_line(lx, y0 + 10, lx + 18, y0 + 10, fill=color,
                          width=max(2, width))
            item = c.create_text(lx + 22, y0 + 10, text=label, anchor="w",
                                 fill=self.AXIS, font=("Segoe UI", 9))
            bbox = c.bbox(item)
            lx = (bbox[2] if bbox else lx + 80) + 14


# --------------------------------------------------------------------------- #
# Experiment export: merge FrED's table with the camera data
# --------------------------------------------------------------------------- #
def _parse_num(cell):
    c = cell.strip()
    if c == "":
        return None
    try:
        return float(c.replace(",", "."))
    except ValueError:
        return cell


def parse_fred_csv(text):
    """FrED's semicolon / comma-decimal table -> (header, rows of values)."""
    reader = csv.reader(io.StringIO(text), delimiter=";")
    header = next(reader, [])
    rows = []
    for raw in reader:
        if not raw or all(not c.strip() for c in raw):
            continue
        rows.append([_parse_num(c) for c in raw])
    return header, rows


def _fmt_cell(value, as_int=False):
    """One CSV cell: comma decimals (Excel es-MX), blank for missing."""
    if value is None or value == "":
        return ""
    if isinstance(value, str):
        return value
    if as_int:
        return str(int(round(value)))
    return f"{float(value):.4f}".replace(".", ",")


def format_semicolon_csv(header, rows):
    ints = [name in INT_COLUMNS for name in header]
    lines = [";".join(header)]
    for row in rows:
        lines.append(";".join(_fmt_cell(v, ints[i] if i < len(ints) else False)
                              for i, v in enumerate(row)))
    return "\r\n".join(lines) + "\r\n"


def merge_run(fred_header, fred_rows, cam_rows, t0_laptop, steady_rel):
    """Put the camera data onto FrED's rows (see the module docstring, 4.).

    ``cam_rows``: MeasurementLog tuples covering the recording (plus margins).
    ``t0_laptop``: FrED's recording start on the laptop clock (None: no merge).
    ``steady_rel``: steady-state mark in recording time (s), or None.
    Returns (header, rows, camera_header, camera_rows, stats).
    """
    L = MeasurementLog
    unit = cam_rows[-1][L.UNIT] if cam_rows else "mm"
    n = len(cam_rows)
    t_rel = [r[L.T] - t0_laptop for r in cam_rows] if t0_laptop is not None else []
    raw = [r[L.D_PX] * r[L.FACTOR] if r[L.FOUND] else None for r in cam_rows]
    found_idx = [i for i in range(n) if cam_rows[i][L.FOUND]]
    filt = [None] * n
    for i, value in zip(found_idx,
                        centered_filter([raw[i] for i in found_idx])):
        filt[i] = float(value)

    diameter_cols = [f"Diameter ({unit})", f"Diameter raw ({unit})",
                     "Diameter new frame", "Diameter camera frame #"]
    if "Diameter setpoint (mm)" in fred_header:
        insert_at = fred_header.index("Diameter setpoint (mm)")
    else:
        insert_at = len(fred_header)
    header = (fred_header[:insert_at] + diameter_cols
              + fred_header[insert_at:] + ["Steady state"])
    t_col = fred_header.index("Time (s)") if "Time (s)" in fred_header else 0

    rows = []
    j = -1
    prev_frame = None
    first_used = None      # index of the first camera frame placed in a row
    new_frames = rows_with_diameter = 0
    for row in fred_rows:
        tr = row[t_col] if isinstance(row[t_col], float) else None
        cells = [None, None, 0, None]
        if tr is not None and n and t0_laptop is not None:
            while j + 1 < n and t_rel[j + 1] <= tr:
                j += 1
            if j >= 0 and tr - t_rel[j] <= STALE_FRAME_S:
                frame_no = cam_rows[j][L.FRAME]
                if first_used is None:
                    first_used = j
                is_new = 1 if frame_no != prev_frame else 0
                prev_frame = frame_no
                new_frames += is_new
                cells = [filt[j], raw[j], is_new, frame_no]
                if raw[j] is not None:
                    rows_with_diameter += 1
            else:
                prev_frame = None
        steady = 1 if (steady_rel is not None and tr is not None
                       and tr >= steady_rel) else 0
        rows.append(row[:insert_at] + cells + row[insert_at:] + [steady])

    # Full-rate camera sheet: every frame inside the recording window, plus
    # the frame held by the first rows (captured just before t = 0), so every
    # frame # referenced in the main table is listed.
    t_first = 0.0 if first_used is None else min(0.0, t_rel[first_used])
    t_last = None
    for row in reversed(fred_rows):
        if isinstance(row[t_col], float):
            t_last = row[t_col]
            break
    cam_header = ["Time (s)", "Camera frame #", "Fiber detected",
                  f"Diameter ({unit})", f"Diameter raw ({unit})",
                  "Diameter (px)", "Min (px)", "Max (px)", "Std (px)",
                  "Steady state"]
    cam_out = []
    if t0_laptop is not None and t_last is not None:
        for i, r in enumerate(cam_rows):
            if not t_first <= t_rel[i] <= t_last:
                continue
            found = bool(r[L.FOUND])
            steady = 1 if (steady_rel is not None and t_rel[i] >= steady_rel) else 0
            cam_out.append([
                t_rel[i], r[L.FRAME], 1 if found else 0, filt[i], raw[i],
                r[L.D_PX] if found else None,
                r[L.MIN_PX] if found else None,
                r[L.MAX_PX] if found else None,
                r[L.STD_PX] if found else None,
                steady,
            ])
    duration = t_last or 0.0
    stats = {
        "unit": unit,
        "frames": len(cam_out),
        "fps": (len(cam_out) / duration) if duration > 0 else 0.0,
        "detected": sum(1 for r in cam_out if r[2]),
        "rows": len(rows),
        "rows_with_diameter": rows_with_diameter,
        "new_frame_rows": new_frames,
        "duration": duration,
    }
    return header, rows, cam_header, cam_out, stats


def write_run_xlsx(path, main_header, main_rows, cam_header, cam_rows,
                   info_rows, progress=None):
    """Formatted Excel copy of the experiment. Returns (ok, message).

    Sheets: "FrED Experiment" (the merged table + native charts), "Camera
    (full rate)" (every camera frame) and "Run info" (timing / sync / params).
    Headers are bold white on a colour per subsystem; row 1 is frozen.
    """
    def report(frac, text):
        if progress is not None:
            progress(frac, text)

    try:
        from openpyxl import Workbook
        from openpyxl.chart import Reference, ScatterChart, Series
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        return False, "openpyxl not installed (pip install openpyxl)"

    def header_color(name):
        n = name.lower()
        if n.startswith("time"):
            return "595959"        # gray
        if n.startswith("temp"):
            return "C0504D"        # red
        if n.startswith("diameter") or n.startswith("fiber"):
            return "4472C4"        # blue
        if n.startswith(("min", "max", "std", "camera")):
            return "5B9BD5"        # light blue (camera details)
        if n.startswith("fan"):
            return "31859C"        # teal
        if n.startswith("extruder"):
            return "7030A0"        # purple
        if n.startswith("spooler"):
            return "548235"        # green
        if n.startswith("steady"):
            return "ED7D31"        # orange
        return "444444"

    def style_header(ws, header):
        for idx, name in enumerate(header, start=1):
            cell = ws.cell(row=1, column=idx)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor=header_color(name))
            cell.alignment = Alignment(horizontal="center",
                                       vertical="center", wrap_text=True)
            ws.column_dimensions[get_column_letter(idx)].width = \
                max(11, min(20, len(name) + 2))
        ws.row_dimensions[1].height = 30
        ws.freeze_panes = "A2"

    def fill(ws, header, rows, frac0, frac1, label):
        ws.append(header)
        total = max(1, len(rows))
        for i, row in enumerate(rows):
            ws.append(["" if v is None else v for v in row])
            if i % 500 == 0:
                report(frac0 + (frac1 - frac0) * i / total,
                       f"Building Excel ({label})... row {i:,} of {total:,}")
        style_header(ws, header)

    try:
        wb = Workbook()
        ws = wb.active
        ws.title = "FrED Experiment"
        fill(ws, main_header, main_rows, 0.05, 0.55, "FrED data")
        ws_cam = None
        if cam_rows:
            ws_cam = wb.create_sheet("Camera (full rate)")
            fill(ws_cam, cam_header, cam_rows, 0.55, 0.75, "camera")
        ws_info = wb.create_sheet("Run info")
        for key, value in info_rows:
            ws_info.append([key, value])
        for cell in ws_info["A"]:
            cell.font = Font(bold=True)
        ws_info.column_dimensions["A"].width = 38
        ws_info.column_dimensions["B"].width = 70

        report(0.8, "Adding charts...")

        def make_chart(title, y_title, series_specs):
            chart = ScatterChart()
            chart.title = title
            chart.style = 13
            chart.x_axis.title = "Time (s)"
            chart.y_axis.title = y_title
            chart.x_axis.delete = False
            chart.y_axis.delete = False
            chart.height = 9
            chart.width = 20
            for sheet, header, n_rows, name in series_specs:
                if sheet is None or name not in header or n_rows < 2:
                    continue
                xref = Reference(sheet, min_col=header.index("Time (s)") + 1,
                                 min_row=2, max_row=n_rows + 1)
                yref = Reference(sheet, min_col=header.index(name) + 1,
                                 min_row=1, max_row=n_rows + 1)
                series = Series(yref, xref, title_from_data=True)
                series.marker.symbol = "none"
                series.smooth = False
                chart.series.append(series)
            return chart if chart.series else None

        n_main, n_cam = len(main_rows), len(cam_rows)
        dia_name = next((h for h in main_header
                         if h.startswith("Diameter (")), "Diameter (mm)")
        raw_name = next((h for h in main_header
                         if h.startswith("Diameter raw (")), "Diameter raw (mm)")
        if ws_cam is not None:
            # The camera sheet has every frame - the true diameter trace.
            dia_specs = [(ws_cam, cam_header, n_cam, cam_header[3]),
                         (ws_cam, cam_header, n_cam, cam_header[4])]
        else:
            dia_specs = [(ws, main_header, n_main, dia_name),
                         (ws, main_header, n_main, raw_name)]
        dia_specs.append((ws, main_header, n_main, "Diameter setpoint (mm)"))
        specs = [
            ("Diameter", dia_name, dia_specs),
            ("Temperature", "Temperature (C)",
             [(ws, main_header, n_main, "Temperature (C)"),
              (ws, main_header, n_main, "Temp setpoint (C)")]),
            ("DC Spooling Motor", "Speed (RPM)",
             [(ws, main_header, n_main, "Spooler RPM"),
              (ws, main_header, n_main, "Spooler setpoint (RPM)")]),
        ]
        anchor_col = get_column_letter(len(main_header) + 2)
        anchor_row = 2
        for title, y_title, series_specs in specs:
            chart = make_chart(title, y_title, series_specs)
            if chart is not None:
                ws.add_chart(chart, f"{anchor_col}{anchor_row}")
                anchor_row += 19   # stack the charts vertically

        report(0.88, "Writing the Excel file (this is the slow part)...")
        wb.save(path)
    except Exception as exc:
        return False, str(exc)
    return True, ""


# --------------------------------------------------------------------------- #
# GUI application
# --------------------------------------------------------------------------- #
class FiberApp:
    GRAPH_WINDOWS = {"10 s": 10.0, "30 s": 30.0, "1 min": 60.0,
                     "2 min": 120.0, "5 min": 300.0, "Whole run": None}
    TICK_MS = 30              # GUI refresh period (display, messages)
    GRAPH_PERIOD_S = 0.1      # graph + reading refresh
    LOG_KEEP_S = 30 * 60      # camera history kept when no run needs it
    WARM_UP_PHASES = ("heating", "extruding", "settle")

    def __init__(self, root):
        self.root = root
        self.root.title(f"FrED Fiber Measure - {APP_VERSION} "
                        "(diameter measured on this laptop)")
        self.root.protocol("WM_DELETE_WINDOW", self.on_close)
        self.root.minsize(960, 620)
        # F11 toggles true fullscreen, Escape leaves it.
        self._fullscreen = False
        self.root.bind("<F11>", self._toggle_fullscreen)
        self.root.bind("<Escape>", self._exit_fullscreen)

        self.calib = Calibration()
        self.detector = FiberDetector()
        self.log = MeasurementLog()
        self.worker = CameraWorker(self.detector, self._on_measurement)
        self.link = FredLink()
        self.sync = ClockSync()

        self.last_result = None
        self.t_origin = time.perf_counter()        # graph origin with no run
        self._wall_offset = time.time() - time.perf_counter()
        # Live (trailing) jitter filter state - measurement thread only.
        self._med_buf = deque(maxlen=FILTER_MEDIAN)
        self._smooth_buf = deque(maxlen=FILTER_MEAN)
        self._filter_unit = None

        # Size the camera feed is scaled to (updated as the window resizes).
        self._feed_w = 640
        self._feed_h = 400

        # Laptop-only recording state (Measure tab)
        self.recording = False
        self.records = []          # list of dicts (one per recorded frame)
        self.rec_start_t = None
        self.frame_counter = 0
        self._rec_target = 0
        self._target_reached = None
        self.save_dir = DATA_DIR   # where Pause & Save writes (default: Data/)

        # FrED experiment state
        self.exp_save_dir = DATA_DIR
        self.run = None            # the latest experiment sent to FrED
        self.pi_phase = "idle"
        self._phase_text = ""
        self._phase_remaining = None   # (seconds left, laptop time received)
        self._retrieve = None          # state of a running Retrieve Data flow
        self._ping_id = 0
        self._last_ping = 0.0
        self._last_graph = 0.0
        self._last_slow = 0.0
        self._last_trim = time.perf_counter()

        self._build_ui()
        self._set_initial_geometry()
        self._update_calib_label()
        self.open_camera()
        self._tick()

    def _set_initial_geometry(self):
        """Open at a size where everything fits with a little free space."""
        self.root.update_idletasks()
        bbox = self.panel_canvas.bbox("all")
        panel_w = (bbox[2] - bbox[0]) if bbox else 480
        panel_h = (bbox[3] - bbox[1]) if bbox else 800

        # Width: control panel + a comfortable video column; height: tall enough
        # for the whole panel plus the status bar and some breathing room.
        want_w = panel_w + 800 + 40
        want_h = panel_h + 90

        # Never exceed the available screen.
        max_w = int(self.root.winfo_screenwidth() * 0.96)
        max_h = int(self.root.winfo_screenheight() * 0.92)
        win_w, win_h = min(want_w, max_w), min(want_h, max_h)

        # Centre the window on screen.
        x = max(0, (self.root.winfo_screenwidth() - win_w) // 2)
        y = max(0, (self.root.winfo_screenheight() - win_h) // 3)
        self.root.geometry(f"{win_w}x{win_h}+{x}+{y}")

    # ------------------------------------------------------------------ #
    # UI construction
    # ------------------------------------------------------------------ #
    def _apply_style(self):
        """Enlarge the fonts/widgets so the control panel is bigger and clearer."""
        # Bump the named fonts every ttk widget inherits from (works on all themes).
        for name, size in (("TkDefaultFont", 11), ("TkTextFont", 11),
                           ("TkMenuFont", 11)):
            try:
                tkfont.nametofont(name).configure(size=size)
            except tk.TclError:
                pass
        style = ttk.Style()
        style.configure(".", font=("Segoe UI", 11))
        style.configure("TButton", font=("Segoe UI", 11), padding=5)
        style.configure("TCheckbutton", font=("Segoe UI", 11))
        style.configure("Toolbutton", font=("Segoe UI", 11), padding=(10, 3))
        style.configure("TLabelframe.Label", font=("Segoe UI", 12, "bold"))
        # Larger font for dropdown lists.
        self.root.option_add("*TCombobox*Listbox.font", ("Segoe UI", 11))

    @staticmethod
    def _big_button(parent, text, command, bg, active_bg):
        return tk.Button(parent, text=text, command=command, bg=bg, fg="white",
                         activebackground=active_bg, activeforeground="white",
                         font=("Segoe UI", 11, "bold"), relief=tk.RAISED,
                         bd=2, padx=12, pady=5, cursor="hand2")

    def _build_ui(self):
        self._apply_style()

        # --- Status bar (packed first so it always stays visible) --------- #
        self.status_var = tk.StringVar(value="Ready.")
        ttk.Label(
            self.root, textvariable=self.status_var, relief=tk.SUNKEN,
            anchor="w", padding=4,
        ).pack(fill=tk.X, side=tk.BOTTOM)

        main = ttk.Frame(self.root, padding=(8, 5))
        main.pack(fill=tk.BOTH, expand=True)
        # Video column expands with the window; the control column keeps its width.
        main.columnconfigure(0, weight=1)
        main.columnconfigure(1, weight=0)
        main.rowconfigure(0, weight=1)

        # --- Left: switchable camera feed on top, diameter graph below --- #
        video_col = ttk.Frame(main)
        video_col.grid(row=0, column=0, sticky="nsew")
        video_col.columnconfigure(0, weight=1)
        video_col.rowconfigure(1, weight=3)   # camera feed
        video_col.rowconfigure(3, weight=2)   # diameter graph

        feed_bar = ttk.Frame(video_col)
        feed_bar.grid(row=0, column=0, sticky="we")
        self.view_var = tk.StringVar(value="live")
        ttk.Radiobutton(feed_bar, text="Live feed", value="live",
                        variable=self.view_var, style="Toolbutton").pack(
            side=tk.LEFT)
        ttk.Radiobutton(feed_bar, text="Mask (processed)", value="mask",
                        variable=self.view_var, style="Toolbutton").pack(
            side=tk.LEFT, padx=(4, 0))
        self.fps_var = tk.StringVar(value="Camera: starting...")
        ttk.Label(feed_bar, textvariable=self.fps_var,
                  foreground="#555").pack(side=tk.RIGHT)

        feed_area = ttk.Frame(video_col)
        feed_area.grid(row=1, column=0, sticky="nsew", pady=(2, 4))
        feed_area.bind("<Configure>", self._on_feed_resize)
        self.feed_label = ttk.Label(feed_area, anchor="center")
        # place() so the image's size can never push the layout around.
        self.feed_label.place(relx=0, rely=0, relwidth=1, relheight=1)

        graph_bar = ttk.Frame(video_col)
        graph_bar.grid(row=2, column=0, sticky="we")
        ttk.Label(graph_bar, text="Diameter",
                  font=("Segoe UI", 11, "bold")).pack(side=tk.LEFT)
        self.graph_value_var = tk.StringVar(value="")
        ttk.Label(graph_bar, textvariable=self.graph_value_var).pack(
            side=tk.LEFT, padx=10)
        self.graph_window_var = tk.StringVar(value="30 s")
        ttk.Combobox(graph_bar, textvariable=self.graph_window_var, width=10,
                     state="readonly", values=list(self.GRAPH_WINDOWS)).pack(
            side=tk.RIGHT)
        ttk.Label(graph_bar, text="Show last:").pack(side=tk.RIGHT, padx=4)

        self.plot = LivePlot(video_col)
        self.plot.canvas.grid(row=3, column=0, sticky="nsew", pady=(2, 0))

        # Run bar: the buttons the operator needs while watching the fiber.
        run_bar = ttk.Frame(video_col)
        run_bar.grid(row=4, column=0, sticky="we", pady=(6, 0))
        self._big_button(run_bar, "START RECORDING NOW",
                         self.start_recording_now, "#2e7d32", "#1b5e20").pack(
            side=tk.LEFT)
        self._big_button(run_bar, "MARK STEADY STATE",
                         self.mark_steady_state, "#e65100", "#bf360c").pack(
            side=tk.LEFT, padx=(8, 0))
        self._big_button(run_bar, "ABORT", self.abort_experiment,
                         "#b71c1c", "#7f0000").pack(side=tk.RIGHT)
        self.run_status_var = tk.StringVar(value="FrED: not connected")
        ttk.Label(video_col, textvariable=self.run_status_var,
                  font=("Segoe UI", 10, "bold"), wraplength=900).grid(
            row=5, column=0, sticky="w", pady=(4, 0))

        # --- Right side: a notebook with the controls and the experiment tab - #
        self.notebook = ttk.Notebook(main)
        self.notebook.grid(row=0, column=1, sticky="ns", padx=(8, 0))
        tab_measure = ttk.Frame(self.notebook)
        self.notebook.add(tab_measure, text="Measure & Connect")
        self._experiment_tab = ttk.Frame(self.notebook)
        self.notebook.add(self._experiment_tab, text="Experiment (FrED)")

        # --- Control panel: sizes to its content; scrollbar only when needed - #
        panel_container = ttk.Frame(tab_measure)
        panel_container.pack(fill=tk.BOTH, expand=True)
        panel_container.rowconfigure(0, weight=1)
        panel_container.columnconfigure(0, weight=1)

        self.panel_canvas = tk.Canvas(panel_container, borderwidth=0,
                                      highlightthickness=0)
        self.panel_vsb = ttk.Scrollbar(panel_container, orient="vertical",
                                       command=self.panel_canvas.yview)
        self.panel_canvas.configure(yscrollcommand=self.panel_vsb.set)
        self.panel_canvas.grid(row=0, column=0, sticky="nsew")
        self.panel_vsb.grid(row=0, column=1, sticky="ns")
        self.panel_vsb.grid_remove()   # hidden until the content is taller than the view

        panel = ttk.Frame(self.panel_canvas, padding=(12, 0))
        self.panel_canvas.create_window((0, 0), window=panel, anchor="nw")
        panel.bind("<Configure>", lambda e: self._update_panel_scroll())
        self.panel_canvas.bind("<Configure>", lambda e: self._update_panel_scroll())

        def _on_wheel(event):
            if self.panel_vsb.winfo_ismapped():   # only scroll when scrollbar is active
                self.panel_canvas.yview_scroll(int(-event.delta / 120), "units")
        self.panel_canvas.bind(
            "<Enter>", lambda e: self.panel_canvas.bind_all("<MouseWheel>", _on_wheel))
        self.panel_canvas.bind(
            "<Leave>", lambda e: self.panel_canvas.unbind_all("<MouseWheel>"))

        # --- Camera ---
        cam_box = ttk.LabelFrame(panel, text="Camera", padding=(8, 5))
        cam_box.pack(fill=tk.X, pady=3)
        ttk.Label(cam_box, text="Index:").grid(row=0, column=0, sticky="w")
        self.cam_index_var = tk.IntVar(value=0)
        ttk.Spinbox(
            cam_box, from_=0, to=10, width=5, textvariable=self.cam_index_var
        ).grid(row=0, column=1, padx=4, sticky="w")
        ttk.Button(cam_box, text="Reconnect", command=self.reconnect).grid(
            row=0, column=2, padx=4
        )
        ttk.Label(cam_box, text="Resolution:").grid(row=1, column=0, sticky="w")
        self.cam_res_var = tk.StringVar(value="Camera default")
        ttk.Combobox(cam_box, textvariable=self.cam_res_var, width=15,
                     state="readonly", values=list(CAMERA_RESOLUTIONS)).grid(
            row=1, column=1, columnspan=2, padx=4, pady=2, sticky="w")
        self.cam_mjpg_var = tk.BooleanVar(value=False)
        ttk.Checkbutton(cam_box, text="High-FPS mode (MJPG)",
                        variable=self.cam_mjpg_var).grid(
            row=2, column=0, columnspan=3, sticky="w")
        ttk.Label(
            cam_box, foreground="#555", wraplength=330,
            text="Every frame the camera delivers is measured. High-FPS mode "
                 "asks for compressed frames: faster if the camera supports "
                 "it, but compression can add edge noise. Press Reconnect to "
                 "apply.").grid(row=3, column=0, columnspan=3, sticky="w")

        # --- Live reading ---
        read_box = ttk.LabelFrame(panel, text="Live measurement", padding=(8, 5))
        read_box.pack(fill=tk.X, pady=3)
        self.reading_var = tk.StringVar(value="-- ")
        ttk.Label(
            read_box, textvariable=self.reading_var,
            font=("Segoe UI", 22, "bold"),
        ).pack(anchor="w")
        self.detail_var = tk.StringVar(value="No fiber detected")
        ttk.Label(read_box, textvariable=self.detail_var).pack(anchor="w")

        # --- Calibration ---
        cal_box = ttk.LabelFrame(panel, text="Calibration", padding=(8, 5))
        cal_box.pack(fill=tk.X, pady=3)
        self.calib_var = tk.StringVar()
        ttk.Label(cal_box, textvariable=self.calib_var).pack(anchor="w")
        btns = ttk.Frame(cal_box)
        btns.pack(anchor="w", pady=(4, 0))
        ttk.Button(
            btns, text="Calibrate (reference)", command=self.calibrate_reference
        ).grid(row=0, column=0, padx=2)
        ttk.Button(
            btns, text="Enter factor", command=self.calibrate_manual
        ).grid(row=0, column=1, padx=2)
        ttk.Button(btns, text="Clear", command=self.clear_calibration).grid(
            row=0, column=2, padx=2
        )

        # --- Detection parameters ---
        par_box = ttk.LabelFrame(panel, text="Detection parameters", padding=(8, 5))
        par_box.pack(fill=tk.X, pady=3)
        par_box.columnconfigure(1, weight=1)

        self.otsu_var = tk.BooleanVar(value=self.detector.use_otsu)
        ttk.Checkbutton(
            par_box, text="Auto threshold (Otsu)", variable=self.otsu_var,
            command=self._sync_params,
        ).grid(row=0, column=0, columnspan=2, sticky="w")

        ttk.Label(par_box, text="Threshold").grid(row=1, column=0, sticky="w")
        self.thresh_var = tk.IntVar(value=self.detector.manual_thresh)
        ttk.Scale(
            par_box, from_=0, to=255, variable=self.thresh_var,
            command=lambda *_: self._sync_params(), length=200,
        ).grid(row=1, column=1, sticky="we")

        ttk.Label(par_box, text="Blur").grid(row=2, column=0, sticky="w")
        self.blur_var = tk.IntVar(value=self.detector.blur_ksize)
        ttk.Scale(
            par_box, from_=1, to=21, variable=self.blur_var,
            command=lambda *_: self._sync_params(), length=200,
        ).grid(row=2, column=1, sticky="we")

        ttk.Label(par_box, text="Min area").grid(row=3, column=0, sticky="w")
        self.area_var = tk.IntVar(value=self.detector.min_area)
        ttk.Scale(
            par_box, from_=50, to=5000, variable=self.area_var,
            command=lambda *_: self._sync_params(), length=200,
        ).grid(row=3, column=1, sticky="we")
        ttk.Button(
            par_box, text="Reset to defaults", command=self.reset_parameters
        ).grid(row=4, column=0, columnspan=2, sticky="w", pady=(6, 0))
        ttk.Label(
            par_box,
            text="Switch the feed to 'Mask (processed)' to see the effect of "
                 "these settings.",
            wraplength=320, foreground="#555",
        ).grid(row=5, column=0, columnspan=2, sticky="w", pady=(4, 0))

        # --- Laptop-only recording ---
        exp_box = ttk.LabelFrame(panel, text="Recording (this laptop only)",
                                 padding=(8, 5))
        exp_box.pack(fill=tk.X, pady=3)
        exp_box.columnconfigure(1, weight=1)

        # Custom file name for the saved data.
        ttk.Label(exp_box, text="File name:").grid(row=0, column=0, sticky="w")
        self.filename_var = tk.StringVar(value="experiment")
        ttk.Entry(exp_box, textvariable=self.filename_var).grid(
            row=0, column=1, columnspan=2, sticky="we", padx=2, pady=2
        )

        # Save folder (defaults to Data/, can be redirected via the file manager).
        ttk.Label(exp_box, text="Save folder:").grid(row=1, column=0, sticky="w")
        self.save_dir_var = tk.StringVar(value=self.save_dir)
        ttk.Label(exp_box, textvariable=self.save_dir_var, foreground="#555",
                  wraplength=210).grid(row=1, column=1, sticky="w", padx=2)
        ttk.Button(exp_box, text="Change...", command=self.change_save_folder).grid(
            row=1, column=2, padx=2
        )

        # Number of samples to record. 0 = unlimited; any positive number
        # auto-stops recording once that many samples are captured.
        ttk.Label(exp_box, text="Samples to record:").grid(
            row=2, column=0, sticky="w"
        )
        self.sample_target_var = tk.IntVar(value=0)
        ttk.Spinbox(
            exp_box, from_=0, to=1000000, increment=10,
            textvariable=self.sample_target_var, width=10,
        ).grid(row=2, column=1, sticky="w", padx=2, pady=2)
        ttk.Label(exp_box, text="(0 = unlimited)", foreground="#555").grid(
            row=2, column=2, sticky="w"
        )

        # Recording / saving buttons.
        self.start_btn = ttk.Button(
            exp_box, text="Start", command=self.start_recording
        )
        self.start_btn.grid(row=3, column=0, padx=2, pady=(6, 0), sticky="we")
        self.pause_btn = ttk.Button(
            exp_box, text="Pause & Save", command=self.pause_recording,
            state=tk.DISABLED,
        )
        self.pause_btn.grid(row=3, column=1, padx=2, pady=(6, 0), sticky="we")
        ttk.Button(exp_box, text="Save As...", command=self.save_as).grid(
            row=3, column=2, padx=2, pady=(6, 0), sticky="we"
        )

        self.rec_status_var = tk.StringVar(value="Idle - 0 samples")
        ttk.Label(exp_box, textvariable=self.rec_status_var).grid(
            row=4, column=0, columnspan=3, sticky="w", pady=(4, 0)
        )
        ttk.Label(
            exp_box,
            text="Saved as both .csv and .xlsx with the same name.",
            foreground="#555", wraplength=300,
        ).grid(row=5, column=0, columnspan=3, sticky="w")

        # --- Connection to the FrED Pi (over WiFi) ---
        link_box = ttk.LabelFrame(
            panel, text="FrED connection (WiFi)", padding=(8, 5)
        )
        link_box.pack(fill=tk.X, pady=3)
        link_box.columnconfigure(1, weight=1)

        # Reminder of which hotspot to join first.
        ttk.Label(
            link_box,
            text=(f"1. Join Wi-Fi  '{PI_HOTSPOT_SSID}'  "
                  f"(password: {PI_HOTSPOT_PASSWORD})"),
            foreground="#225522", wraplength=320,
        ).grid(row=0, column=0, columnspan=3, sticky="w")
        ttk.Label(
            link_box, text="2. Enter the Pi address shown on its screen:",
            foreground="#555", wraplength=320,
        ).grid(row=1, column=0, columnspan=3, sticky="w", pady=(2, 4))

        ttk.Label(link_box, text="Pi IP:").grid(row=2, column=0, sticky="w")
        self.host_var = tk.StringVar(value=PI_DEFAULT_IP)
        ttk.Entry(link_box, textvariable=self.host_var, width=16).grid(
            row=2, column=1, padx=4, sticky="we"
        )

        ttk.Label(link_box, text="Port:").grid(row=3, column=0, sticky="w")
        self.port_var = tk.StringVar(value=str(PI_DEFAULT_PORT))
        ttk.Entry(link_box, textvariable=self.port_var, width=8).grid(
            row=3, column=1, padx=4, sticky="w"
        )

        self.connect_btn = ttk.Button(
            link_box, text="Connect", command=self.toggle_connection
        )
        self.connect_btn.grid(row=4, column=0, columnspan=2, pady=(6, 0),
                              padx=2, sticky="we")

        self.link_status_var = tk.StringVar(value="Not connected")
        ttk.Label(link_box, textvariable=self.link_status_var,
                  wraplength=320).grid(
            row=5, column=0, columnspan=3, sticky="w", pady=(4, 0)
        )
        ttk.Label(
            link_box, foreground="#555", wraplength=320,
            text="The diameter stays on this laptop: the link only carries "
                 "experiment commands and a tiny clock-sync ping every "
                 f"{SYNC_INTERVAL_S:.0f} s.").grid(
            row=6, column=0, columnspan=3, sticky="w", pady=(2, 0))

        # --- Experiment tab (send a whole run to FrED) ---
        self._build_experiment_tab(self._experiment_tab)

    # ------------------------------------------------------------------ #
    # Camera handling
    # ------------------------------------------------------------------ #
    def open_camera(self):
        try:
            index = int(self.cam_index_var.get())
        except (tk.TclError, ValueError):
            index = 0
        resolution = CAMERA_RESOLUTIONS.get(self.cam_res_var.get())
        if self.worker.start(index, resolution, self.cam_mjpg_var.get()):
            fmt = f" {self.worker.fourcc}" if self.worker.fourcc else ""
            self.status_var.set(
                f"Camera {index} connected ({self.worker.width}x"
                f"{self.worker.height}{fmt}).")
        else:
            self.fps_var.set("Camera: not available")
            self.status_var.set(
                f"Could not open camera {index}. "
                "Check the index and reconnect."
            )

    def reconnect(self):
        if self.recording:
            messagebox.showwarning(
                "Recording", "Stop the recording before switching cameras."
            )
            return
        self.open_camera()

    # ------------------------------------------------------------------ #
    # Parameter syncing
    # ------------------------------------------------------------------ #
    def _sync_params(self):
        self.detector.use_otsu = self.otsu_var.get()
        self.detector.manual_thresh = int(self.thresh_var.get())
        self.detector.blur_ksize = int(self.blur_var.get())
        self.detector.min_area = int(self.area_var.get())

    def reset_parameters(self):
        """Restore the preset detection parameters."""
        d = FiberDetector.DEFAULTS
        self.otsu_var.set(d["use_otsu"])
        self.thresh_var.set(d["manual_thresh"])
        self.blur_var.set(d["blur_ksize"])
        self.area_var.set(d["min_area"])
        self._sync_params()
        self.status_var.set("Detection parameters reset to defaults.")

    # ------------------------------------------------------------------ #
    # Measurement (runs on the camera worker's measurement thread)
    # ------------------------------------------------------------------ #
    def _on_measurement(self, frame_no, t, result):
        """Store one measured frame. NO Tk calls in here (worker thread)."""
        factor, unit = self.calib.display_scale()
        found = result["found"]
        d_px = result["diameter_px"]
        filtered = None
        if found:
            if unit != self._filter_unit:     # calibration changed
                self._med_buf.clear()
                self._smooth_buf.clear()
                self._filter_unit = unit
            self._med_buf.append(d_px * factor)
            self._smooth_buf.append(statistics.median(self._med_buf))
            filtered = sum(self._smooth_buf) / len(self._smooth_buf)
        self.log.append((t, frame_no, found, d_px, result["min_px"],
                         result["max_px"], result["std_px"], factor, unit,
                         filtered))
        self.last_result = result
        if self.recording and found:
            self._record(t, result)

    # ------------------------------------------------------------------ #
    # Main loop (Tk thread)
    # ------------------------------------------------------------------ #
    def _tick(self):
        now = time.perf_counter()
        # Handle any messages FrED sent back (status / data / sync).
        self._drain_pi_messages()

        view = self.worker.take_view()
        if view is not None:
            self._show_frame(view)
        self.worker.request_view(self.view_var.get())

        if self._target_reached is not None:
            target, self._target_reached = self._target_reached, None
            self._finish_target_recording(target)

        if now - self._last_graph >= self.GRAPH_PERIOD_S:
            self._last_graph = now
            self._update_reading()
            self._update_graph(now)

        if now - self._last_slow >= 1.0:
            self._last_slow = now
            self._update_slow_status(now)

        if self.link.is_open and now - self._last_ping >= SYNC_INTERVAL_S:
            self._last_ping = now
            self._send_ping()

        self.root.after(self.TICK_MS, self._tick)

    def _send_ping(self):
        self._ping_id += 1
        if not self.link.ping(self._ping_id):
            self._mark_link_lost()

    def _sync_burst(self, count=10, spacing_ms=60):
        """A quick series of pings (right after connecting)."""
        for i in range(count):
            self.root.after(i * spacing_ms,
                            lambda: self.link.is_open and self._send_ping())

    def _drain_pi_messages(self):
        """Process Pi -> laptop messages on the Tk main thread (safe for UI)."""
        while True:
            try:
                msg = self.link.rx_queue.get_nowait()
            except queue.Empty:
                break
            self._handle_pi_message(msg)

    def _update_slow_status(self, now):
        """Once a second: camera rate, link/sync, run status, housekeeping."""
        w = self.worker
        if w.running:
            fmt = f" {w.fourcc}" if w.fourcc else ""
            self.fps_var.set(
                f"Camera {w.width}x{w.height}{fmt}: {w.capture_fps:.1f} fps  |  "
                f"measured {w.measure_fps:.1f} fps  |  dropped {w.dropped}")
        if w.error:
            self.status_var.set(w.error + " Press Reconnect.")
            w.error = None

        if self.link.is_open:
            est = self.sync.estimate()
            sync_text = (f"clock sync ±{est['err'] * 1000:.1f} ms "
                         f"({est['n']} pings)" if est else "clock sync: waiting")
            self.link_status_var.set(
                f"Connected to {self.link.endpoint}  |  {sync_text}")
        # Keep the live t = 0 estimate current as more sync samples arrive.
        if self.run is not None and self.run.get("t0_pi") is not None:
            self.run["t0_laptop"] = self._pi_to_laptop(
                self.run["t0_pi"], self.run.get("boot"),
                self.run.get("t0_arrival"))[0]
        self._update_run_status()

        if self.recording:
            n = len(self.records)
            suffix = f"/{self._rec_target}" if self._rec_target else ""
            self.rec_status_var.set(f"Recording - {n}{suffix} samples")

        if now - self._last_trim >= 60.0:
            self._last_trim = now
            keep_from = now - self.LOG_KEEP_S
            if self.run is not None:
                keep_from = min(keep_from, self.run["sent_t"] - 60.0)
            self.log.trim_before(keep_from)

    # ------------------------------------------------------------------ #
    # Display sizing / fullscreen
    # ------------------------------------------------------------------ #
    def _on_feed_resize(self, event):
        """Track the feed area size so the video scales to fill it."""
        self._feed_w = max(160, event.width - 4)
        self._feed_h = max(120, event.height - 4)

    def _update_panel_scroll(self):
        """Size the control canvas to its content; show the scrollbar only if needed."""
        bbox = self.panel_canvas.bbox("all")
        if bbox is None:
            return
        content_w = bbox[2] - bbox[0]
        content_h = bbox[3] - bbox[1]
        # Keep the canvas exactly as wide as the controls so nothing is clipped.
        if self.panel_canvas.winfo_width() != content_w:
            self.panel_canvas.configure(width=content_w)
        self.panel_canvas.configure(scrollregion=bbox)
        if content_h > self.panel_canvas.winfo_height() + 1:
            self.panel_vsb.grid()
        else:
            self.panel_vsb.grid_remove()
            self.panel_canvas.yview_moveto(0)

    def _toggle_fullscreen(self, _event=None):
        self._fullscreen = not self._fullscreen
        self.root.attributes("-fullscreen", self._fullscreen)

    def _exit_fullscreen(self, _event=None):
        if self._fullscreen:
            self._fullscreen = False
            self.root.attributes("-fullscreen", False)

    def _show_frame(self, bgr):
        """Scale a BGR frame to the feed area and show it."""
        rgb = cv2.cvtColor(bgr, cv2.COLOR_BGR2RGB)
        h, w = rgb.shape[:2]
        scale = min(self._feed_w / w, self._feed_h / h)
        if scale <= 0:
            scale = 1.0
        new_w, new_h = max(1, int(w * scale)), max(1, int(h * scale))
        interp = cv2.INTER_AREA if scale < 1 else cv2.INTER_LINEAR
        rgb = cv2.resize(rgb, (new_w, new_h), interpolation=interp)
        imgtk = ImageTk.PhotoImage(image=Image.fromarray(rgb))
        self.feed_label.imgtk = imgtk  # keep a reference to avoid GC
        self.feed_label.configure(image=imgtk)

    def _update_reading(self):
        row = self.log.last()
        if row is None or not row[MeasurementLog.FOUND]:
            self.reading_var.set("--")
            self.detail_var.set("No fiber detected")
            return
        L = MeasurementLog
        d_px, factor, unit = row[L.D_PX], row[L.FACTOR], row[L.UNIT]
        if unit == "px":
            self.reading_var.set(f"{d_px:.1f} px")
            self.detail_var.set(
                f"min {row[L.MIN_PX]:.1f}  max {row[L.MAX_PX]:.1f}  "
                f"std {row[L.STD_PX]:.1f} px  (not calibrated)")
        else:
            self.reading_var.set(f"{d_px * factor:.4f} {unit}")
            self.detail_var.set(
                f"{d_px:.1f} px | min {row[L.MIN_PX] * factor:.4f} "
                f"max {row[L.MAX_PX] * factor:.4f} "
                f"std {row[L.STD_PX] * factor:.4f} {unit}")

    def _graph_origin(self):
        """(laptop time shown as x = 0, x-axis label)."""
        if self.run is not None and self.run.get("t0_laptop") is not None:
            return self.run["t0_laptop"], "Time since FrED recording start (s)"
        return self.t_origin, "Time (s)"

    def _update_graph(self, now):
        L = MeasurementLog
        origin, x_label = self._graph_origin()
        window = self.GRAPH_WINDOWS.get(self.graph_window_var.get(), 30.0)
        if window is None:
            if self.run is not None:
                t_start = min(self.run["sent_t"], origin)
            else:
                t_start = self.log.first_time() or now
        else:
            t_start = now - window
        rows = self.log.since(t_start)
        # Thin out long windows: ~2 points per pixel column is plenty.
        max_points = max(200, 2 * self.plot.canvas.winfo_width())
        step = max(1, len(rows) // max_points)
        xs_raw, ys_raw, xs_f, ys_f = [], [], [], []
        for r in rows[::step]:
            if not r[L.FOUND]:
                continue
            x = r[L.T] - origin
            xs_raw.append(x)
            ys_raw.append(r[L.D_PX] * r[L.FACTOR])
            if r[L.FILT] is not None:
                xs_f.append(x)
                ys_f.append(r[L.FILT])
        last = rows[-1] if rows else self.log.last()
        unit = last[L.UNIT] if last else self.calib.display_scale()[1]

        hlines = []
        target = self._target_diameter()
        if target is not None and unit == "mm":
            hlines.append((target, "#d32f2f", f"Target {target:g} mm"))
        vlines = []
        if self.run is not None:
            if self.run.get("t0_laptop") is not None:
                vlines.append((self.run["t0_laptop"] - origin, "#2e7d32",
                               "Recording start"))
            if self.run.get("steady_t") is not None:
                vlines.append((self.run["steady_t"] - origin, "#e65100",
                               "Steady state"))
        self.plot.draw(
            t_start - origin, now - origin,
            [(xs_raw, ys_raw, "#9ec5ea", 1, "Raw (every frame)"),
             (xs_f, ys_f, "#0d47a1", 2, "Filtered")],
            x_label, f"Diameter ({unit})", hlines, vlines)
        if last is not None and last[L.FOUND] and last[L.FILT] is not None:
            decimals = 1 if unit == "px" else 4
            self.graph_value_var.set(
                f"{last[L.FILT]:.{decimals}f} {unit} (filtered)")
        else:
            self.graph_value_var.set("no fiber detected")

    def _target_diameter(self):
        try:
            value = float(self.exp_target_diameter_var.get())
        except (tk.TclError, ValueError, AttributeError):
            return None
        return value if value > 0 else None

    # ------------------------------------------------------------------ #
    # Calibration actions
    # ------------------------------------------------------------------ #
    def _update_calib_label(self):
        if self.calib.is_calibrated:
            self.calib_var.set(
                f"Calibrated: {self.calib.units_per_pixel:.6f} "
                f"{self.calib.units}/px"
            )
        else:
            self.calib_var.set("Not calibrated (measuring in pixels)")

    def calibrate_reference(self):
        if self.last_result is None or not self.last_result["found"]:
            messagebox.showwarning(
                "Calibration",
                "No fiber detected. Place a reference object of known size in "
                "view, make sure it is detected, then calibrate.",
            )
            return
        measured_px = self.last_result["diameter_px"]
        units = simpledialog.askstring(
            "Calibration", "Units (e.g. mm, um):", initialvalue=self.calib.units,
            parent=self.root,
        )
        if not units:
            return
        known = simpledialog.askfloat(
            "Calibration",
            f"The detected object measures {measured_px:.1f} px.\n"
            f"Enter its true diameter in {units}:",
            parent=self.root, minvalue=0.0,
        )
        if known is None or known <= 0:
            return
        self.calib.set_from_reference(measured_px, known, units)
        self._update_calib_label()
        self.status_var.set(
            f"Calibrated: {self.calib.units_per_pixel:.6f} {units}/px"
        )

    def calibrate_manual(self):
        units = simpledialog.askstring(
            "Calibration", "Units (e.g. mm, um):", initialvalue=self.calib.units,
            parent=self.root,
        )
        if not units:
            return
        factor = simpledialog.askfloat(
            "Calibration", f"Enter {units} per pixel:", parent=self.root,
            minvalue=0.0,
        )
        if factor is None or factor <= 0:
            return
        self.calib.set_factor(factor, units)
        self._update_calib_label()
        self.status_var.set(f"Calibration factor set: {factor:.6f} {units}/px")

    def clear_calibration(self):
        self.calib.clear()
        self._update_calib_label()
        self.status_var.set("Calibration cleared. Measuring in pixels.")

    # ------------------------------------------------------------------ #
    # Laptop-only recording (Measure tab)
    # ------------------------------------------------------------------ #
    def start_recording(self):
        if self.recording:
            return
        self.records = []
        self.frame_counter = 0
        self._rec_target = self._sample_target()
        self.rec_start_t = time.perf_counter()
        self.recording = True
        self.start_btn.configure(state=tk.DISABLED)
        self.pause_btn.configure(state=tk.NORMAL)
        self.status_var.set("Recording started.")

    def _record(self, t, result):
        """Append one detected frame (measurement thread - no Tk calls)."""
        elapsed = t - self.rec_start_t
        self.frame_counter += 1
        d_px = result["diameter_px"]
        stamp = _dt.datetime.fromtimestamp(t + self._wall_offset)
        row = {
            "timestamp": stamp.strftime("%Y-%m-%d %H:%M:%S.%f")[:-3],
            "elapsed_s": round(elapsed, 4),
            "frame": self.frame_counter,
            "diameter_px": round(d_px, 3),
            "min_px": round(result["min_px"], 3),
            "max_px": round(result["max_px"], 3),
            "std_px": round(result["std_px"], 3),
            "length_px": round(result["length_px"], 3),
            "angle_deg": round(result["angle_deg"], 3),
        }
        if self.calib.is_calibrated:
            row["diameter_real"] = round(self.calib.to_real(d_px), 6)
            row["units"] = self.calib.units
        else:
            row["diameter_real"] = ""
            row["units"] = "px"
        self.records.append(row)
        if self._rec_target and len(self.records) >= self._rec_target:
            self.recording = False
            self._target_reached = self._rec_target   # handled on the Tk thread

    def _sample_target(self):
        """Requested number of samples (0/blank/invalid -> unlimited)."""
        try:
            target = int(self.sample_target_var.get())
        except (tk.TclError, ValueError):
            return 0
        return target if target > 0 else 0

    def _finish_target_recording(self, target):
        """Auto-stop once the requested sample count is reached.

        Recording stops and the buttons return to their normal idle state. The
        data is kept; the user is offered the same save as Pause & Save (and
        can still use Save As... or Start again if they decline)."""
        self.start_btn.configure(state=tk.NORMAL)
        self.pause_btn.configure(state=tk.DISABLED)
        n = len(self.records)
        self.rec_status_var.set(f"Done - {n} samples (target {target})")
        self.status_var.set(f"Reached target of {target} samples.")

        save = messagebox.askyesno(
            "Target reached",
            f"Recorded the requested {target} samples.\n\n"
            f"Save the data (.csv and .xlsx) to:\n{self.save_dir}?",
        )
        if save and self._save_dataset(self.save_dir, self._base_filename()):
            self.records = []
            self.rec_status_var.set("Idle - 0 samples")

    def change_save_folder(self):
        """Pick the folder Pause & Save writes into (Windows file manager)."""
        folder = filedialog.askdirectory(
            title="Choose folder to save experiment data",
            initialdir=self.save_dir,
        )
        if folder:
            self.save_dir = folder
            self.save_dir_var.set(folder)
            self.status_var.set(f"Save folder set to {folder}")

    def _base_filename(self):
        """Sanitised base name (no extension) from the File name box."""
        name = self.filename_var.get().strip()
        if not name:
            name = "experiment_" + _dt.datetime.now().strftime("%Y%m%d_%H%M%S")
        base, ext = os.path.splitext(name)
        if ext.lower() in (".csv", ".xlsx"):
            name = base
        for ch in '<>:"/\\|?*':       # characters illegal in Windows file names
            name = name.replace(ch, "_")
        return name or "experiment"

    def pause_recording(self):
        if not self.recording:
            return
        self.recording = False
        self.start_btn.configure(state=tk.NORMAL)
        self.pause_btn.configure(state=tk.DISABLED)
        n = len(self.records)
        self.rec_status_var.set(f"Paused - {n} samples")

        if n == 0:
            messagebox.showinfo(
                "No data", "No samples were recorded; nothing to save."
            )
            return

        save = messagebox.askyesno(
            "Save recording",
            f"Recording paused with {n} samples.\n\n"
            f"Save the data (.csv and .xlsx) to:\n{self.save_dir}?",
        )
        if save:
            if self._save_dataset(self.save_dir, self._base_filename()):
                self.records = []
                self.rec_status_var.set("Idle - 0 samples")
        else:
            keep = messagebox.askyesno(
                "Discard or continue",
                "Data was NOT saved.\n\n"
                "Yes = discard this data\nNo = keep it (Start again to add more, "
                "Pause again to save, or use Save As...).",
            )
            if keep:
                self.records = []
                self.rec_status_var.set("Idle - 0 samples")
                self.status_var.set("Recorded data discarded.")

    def save_as(self):
        """Save the current data to a location chosen in the Windows file dialog."""
        if self.recording:
            messagebox.showwarning(
                "Recording", "Pause the recording before saving."
            )
            return
        if not self.records:
            messagebox.showinfo("No data", "There is no recorded data to save.")
            return
        path = filedialog.asksaveasfilename(
            title="Save recorded data as",
            initialdir=self.save_dir,
            initialfile=self._base_filename() + ".csv",
            defaultextension=".csv",
            filetypes=[("CSV and Excel", "*.csv"), ("All files", "*.*")],
        )
        if not path:
            return
        folder = os.path.dirname(path)
        base = os.path.splitext(os.path.basename(path))[0]
        self._save_dataset(folder, base)

    def _save_dataset(self, folder, base):
        """Write <base>.csv and <base>.xlsx into ``folder``. Returns True on success."""
        try:
            os.makedirs(folder, exist_ok=True)
        except OSError as exc:
            messagebox.showerror("Save failed", f"Could not create folder:\n{exc}")
            return False

        n = len(self.records)
        csv_path = os.path.join(folder, base + ".csv")
        xlsx_path = os.path.join(folder, base + ".xlsx")

        try:
            with open(csv_path, "w", newline="", encoding="utf-8") as fh:
                writer = csv.DictWriter(fh, fieldnames=EXPORT_FIELDS)
                writer.writeheader()
                for row in self.records:
                    writer.writerow(row)
        except OSError as exc:
            messagebox.showerror("Save failed", f"Could not write CSV:\n{exc}")
            return False

        xlsx_ok, xlsx_msg = self._write_xlsx(xlsx_path)

        saved = csv_path + ("\n" + xlsx_path if xlsx_ok else "")
        note = "" if xlsx_ok else f"\n\nNote: Excel file not written ({xlsx_msg})."
        self.rec_status_var.set(f"Saved {n} samples")
        self.status_var.set(f"Saved {csv_path}")
        messagebox.showinfo(
            "Saved",
            f"Saved {n} samples to:\n{saved}\n\nDiameter: {self._summary()}{note}",
        )
        return True

    def _write_xlsx(self, path):
        """Write the records to an .xlsx file. Returns (ok, message)."""
        try:
            from openpyxl import Workbook
        except ImportError:
            return False, "openpyxl not installed (pip install openpyxl)"
        try:
            wb = Workbook()
            ws = wb.active
            ws.title = "Measurements"
            ws.append(EXPORT_FIELDS)
            for row in self.records:
                ws.append([row.get(f, "") for f in EXPORT_FIELDS])
            wb.save(path)
        except Exception as exc:
            return False, str(exc)
        return True, ""

    def _summary(self):
        d = np.array([r["diameter_px"] for r in self.records], dtype=float)
        if d.size == 0:
            return "no samples"
        if self.calib.is_calibrated:
            d_real = d * self.calib.units_per_pixel
            return (f"mean {d_real.mean():.4f} {self.calib.units}, "
                    f"std {d_real.std():.4f} {self.calib.units}")
        return f"mean {d.mean():.2f} px, std {d.std():.2f} px"

    # ------------------------------------------------------------------ #
    # Connection to the FrED Pi
    # ------------------------------------------------------------------ #
    def toggle_connection(self):
        if self.link.is_open:
            self.link.close()
            self.connect_btn.configure(text="Connect")
            self.link_status_var.set("Not connected")
            self.status_var.set("WiFi link closed.")
            return

        host = self.host_var.get().strip()
        port = self.port_var.get().strip()
        if not host or not port:
            messagebox.showwarning(
                "Connect", "Enter the Pi's IP address and port (shown on the "
                "Pi's screen). Default is 192.168.4.1 : 5005."
            )
            return
        try:
            port_num = int(port)
        except ValueError:
            messagebox.showwarning("Connect", f"'{port}' is not a valid port number.")
            return
        try:
            self.link.open(host, port_num)
        except Exception as exc:
            messagebox.showerror(
                "Connect failed",
                f"Could not connect to {host}:{port_num}\n\n{exc}\n\n"
                f"Make sure the laptop is joined to the '{PI_HOTSPOT_SSID}' "
                "Wi-Fi and the Pi program is running."
            )
            self.link_status_var.set("Connection failed.")
            return
        self.connect_btn.configure(text="Disconnect")
        self.link_status_var.set(f"Connected to {self.link.endpoint}")
        self.status_var.set(f"Connected to {self.link.endpoint}.")
        self._last_ping = time.perf_counter()
        self._sync_burst()

    def _mark_link_lost(self):
        """Close the link and put the connection UI in the 'lost' state."""
        self.link.close()
        self.connect_btn.configure(text="Connect")
        self.link_status_var.set("Link lost - reconnect to resume.")

    def _reset_link(self):
        """Close and re-open the TCP connection to FrED (a fresh stream clears
        any half-received message). Returns True if the link is open again."""
        host, port = self.link.host, self.link.port
        if host is None:
            return False
        try:
            self.link.open(host, port)
        except Exception:
            self._mark_link_lost()
            return False
        self.connect_btn.configure(text="Disconnect")
        self.link_status_var.set(f"Connected to {self.link.endpoint}")
        return True

    def _pi_to_laptop(self, t_pi, boot=None, fallback=None):
        """Pi clock -> laptop clock. Returns (t_laptop, method text)."""
        est = self.sync.estimate(boot)
        if est is not None:
            return (ClockSync.to_laptop(est, t_pi),
                    f"clock sync, ±{est['err'] * 1000:.1f} ms")
        return fallback, "arrival time of FrED's message (no clock sync)"

    # ================================================================== #
    # Experiment tab: configure a run, send it to FrED, retrieve the data
    # ================================================================== #
    def _scrollable(self, parent):
        """Create a vertically scrollable frame inside ``parent``; return it."""
        canvas = tk.Canvas(parent, borderwidth=0, highlightthickness=0)
        vsb = ttk.Scrollbar(parent, orient="vertical", command=canvas.yview)
        canvas.configure(yscrollcommand=vsb.set)
        canvas.pack(side="left", fill="both", expand=True)
        vsb.pack(side="right", fill="y")
        inner = ttk.Frame(canvas, padding=(10, 8))
        canvas.create_window((0, 0), window=inner, anchor="nw")
        inner.bind("<Configure>",
                   lambda e: canvas.configure(scrollregion=canvas.bbox("all")))
        canvas.bind("<Enter>", lambda e: canvas.bind_all(
            "<MouseWheel>",
            lambda ev: canvas.yview_scroll(int(-ev.delta / 120), "units")))
        canvas.bind("<Leave>", lambda e: canvas.unbind_all("<MouseWheel>"))
        return inner

    @staticmethod
    def _exp_row(parent, label, var, row, width=10):
        ttk.Label(parent, text=label).grid(row=row, column=0, sticky="w", pady=2)
        ttk.Entry(parent, textvariable=var, width=width).grid(
            row=row, column=1, sticky="w", padx=4, pady=2)

    def _build_experiment_tab(self, parent):
        panel = self._scrollable(parent)

        ttk.Label(panel, text="Send a full experiment to FrED",
                  font=("Segoe UI", 12, "bold")).pack(anchor="w")
        ttk.Label(panel, wraplength=360, foreground="#555",
                  text="Connect on the 'Measure & Connect' tab first. The "
                       "diameter is measured and recorded on this laptop, so "
                       "keep the camera running and calibrated; it is merged "
                       "with FrED's data when you click Retrieve Data.").pack(
            anchor="w", pady=(0, 6))

        # --- Experiment sequence timing ---
        tbox = ttk.LabelFrame(panel, text="Experiment sequence", padding=(8, 5))
        tbox.pack(fill=tk.X, pady=3)
        self.exp_heating_delay_var = tk.StringVar(value="60")
        self.exp_heat_extrude_time_var = tk.StringVar(value="30")
        self.exp_heat_extrude_speed_var = tk.StringVar(value="1.5")
        self.exp_data_delay_var = tk.StringVar(value="10")
        self.exp_data_time_var = tk.StringVar(value="120")
        self.exp_post_spool_var = tk.StringVar(value="15")
        self.exp_rate_var = tk.StringVar(value="50")
        self._exp_row(tbox, "Heating time (s) - heater only",
                      self.exp_heating_delay_var, 0)
        self._exp_row(tbox, "Heating + extrusion time (s)",
                      self.exp_heat_extrude_time_var, 1)
        self._exp_row(tbox, "Extrusion rate during it (RPM)",
                      self.exp_heat_extrude_speed_var, 2)
        self._exp_row(tbox, "Experiment settle time (s)",
                      self.exp_data_delay_var, 3)
        self._exp_row(tbox, "Data-taking time (s)", self.exp_data_time_var, 4)
        self._exp_row(tbox, "Extra spooling after end (s)",
                      self.exp_post_spool_var, 5)
        self._exp_row(tbox, "FrED sample rate (Hz, 1-100)",
                      self.exp_rate_var, 6)
        ttk.Label(tbox, foreground="#555", wraplength=340,
                  text="Sequence: heat only -> heat + extrude (at the rate "
                       "above) -> settle with all systems on -> record -> "
                       "everything stops except the spooler, which keeps "
                       "coiling fiber for the extra spooling time. START "
                       "RECORDING NOW (under the graph) skips straight to "
                       "recording the moment the fiber drops. The sample "
                       "rate is FrED's control AND data rate.").grid(
            row=7, column=0, columnspan=2, sticky="w", pady=(4, 0))

        # --- Heater ---
        hbox = ttk.LabelFrame(panel, text="Heater", padding=(8, 5))
        hbox.pack(fill=tk.X, pady=3)
        ttk.Label(hbox, text="Mode").grid(row=0, column=0, sticky="w")
        self.exp_heater_mode_var = tk.StringVar(value="closed")
        ttk.Combobox(hbox, textvariable=self.exp_heater_mode_var, width=8,
                     state="readonly", values=["closed", "open"]).grid(
            row=0, column=1, sticky="w", padx=4)
        self.exp_target_temp_var = tk.StringVar(value="95")
        self.exp_temp_kp_var = tk.StringVar(value="1.0")
        self.exp_temp_ki_var = tk.StringVar(value="0.001")
        self.exp_temp_kd_var = tk.StringVar(value="0.05")
        self.exp_heater_pwm_var = tk.StringVar(value="0")
        self._exp_row(hbox, "Target temp (C) [closed]", self.exp_target_temp_var, 1)
        self._exp_row(hbox, "Temp Kp [closed]", self.exp_temp_kp_var, 2)
        self._exp_row(hbox, "Temp Ki [closed]", self.exp_temp_ki_var, 3)
        self._exp_row(hbox, "Temp Kd [closed]", self.exp_temp_kd_var, 4)
        self._exp_row(hbox, "Heater PWM (%) [open]", self.exp_heater_pwm_var, 5)

        # --- Spooler ---
        kp_max, ki_max, kd_max = MOTOR_GAIN_LIMITS
        sbox = ttk.LabelFrame(panel, text="Spooler (DC motor)", padding=(8, 5))
        sbox.pack(fill=tk.X, pady=3)
        ttk.Label(sbox, text="Mode").grid(row=0, column=0, sticky="w")
        self.exp_spooler_mode_var = tk.StringVar(value="closed")
        ttk.Combobox(sbox, textvariable=self.exp_spooler_mode_var, width=8,
                     state="readonly", values=["closed", "open"]).grid(
            row=0, column=1, sticky="w", padx=4)
        self.exp_motor_setpoint_var = tk.StringVar(value="30")
        self.exp_motor_kp_var = tk.StringVar(value="0.5")
        self.exp_motor_ki_var = tk.StringVar(value="0.5")
        self.exp_motor_kd_var = tk.StringVar(value="0.05")
        self.exp_dc_pwm_var = tk.StringVar(value="0")
        self._exp_row(sbox, "Setpoint (RPM) [closed]", self.exp_motor_setpoint_var, 1)
        self._exp_row(sbox, f"Motor Kp [closed, max {kp_max:g}]",
                      self.exp_motor_kp_var, 2)
        self._exp_row(sbox, f"Motor Ki [closed, max {ki_max:g}]",
                      self.exp_motor_ki_var, 3)
        self._exp_row(sbox, f"Motor Kd [closed, max {kd_max:g}]",
                      self.exp_motor_kd_var, 4)
        self._exp_row(sbox, "DC Motor PWM (%) [open]", self.exp_dc_pwm_var, 5)

        # --- Stepper / Fan / Diameter ---
        obox = ttk.LabelFrame(panel, text="Extruder / Fan / Diameter",
                              padding=(8, 5))
        obox.pack(fill=tk.X, pady=3)
        self.exp_extrusion_var = tk.StringVar(value="1.5")
        self.exp_fan_var = tk.StringVar(value="40")
        self.exp_target_diameter_var = tk.StringVar(value="0.35")
        self._exp_row(obox, "Extrusion speed (RPM)", self.exp_extrusion_var, 0)
        self._exp_row(obox, "Fan duty (%)", self.exp_fan_var, 1)
        self._exp_row(obox, "Target diameter (mm)", self.exp_target_diameter_var, 2)

        # --- Save settings ---
        vbox = ttk.LabelFrame(panel, text="Save received data", padding=(8, 5))
        vbox.pack(fill=tk.X, pady=3)
        vbox.columnconfigure(1, weight=1)
        ttk.Label(vbox, text="File name:").grid(row=0, column=0, sticky="w")
        self.exp_name_var = tk.StringVar(value="fred_experiment")
        ttk.Entry(vbox, textvariable=self.exp_name_var).grid(
            row=0, column=1, columnspan=2, sticky="we", padx=2, pady=2)
        ttk.Label(vbox, text="Folder:").grid(row=1, column=0, sticky="w")
        self.exp_save_dir_var = tk.StringVar(value=self.exp_save_dir)
        ttk.Label(vbox, textvariable=self.exp_save_dir_var, foreground="#555",
                  wraplength=210).grid(row=1, column=1, sticky="w", padx=2)
        ttk.Button(vbox, text="Change...", command=self.change_exp_folder).grid(
            row=1, column=2, padx=2)

        # --- Action buttons ---
        btns = ttk.Frame(panel)
        btns.pack(fill=tk.X, pady=(6, 2))
        ttk.Button(btns, text="Send & Start Experiment",
                   command=self.send_experiment).pack(fill=tk.X, pady=2)
        row2 = ttk.Frame(btns)
        row2.pack(fill=tk.X)
        ttk.Button(row2, text="Abort", command=self.abort_experiment).pack(
            side=tk.LEFT, expand=True, fill=tk.X, padx=(0, 2))
        self.exp_retrieve_btn = ttk.Button(
            row2, text="Retrieve Data", command=self.retrieve_data)
        self.exp_retrieve_btn.pack(side=tk.LEFT, expand=True, fill=tk.X, padx=(2, 0))

        self.exp_status_var = tk.StringVar(value="Experiment: idle")
        ttk.Label(panel, textvariable=self.exp_status_var, wraplength=360,
                  font=("Segoe UI", 11, "bold")).pack(anchor="w", pady=(6, 0))

    def change_exp_folder(self):
        folder = filedialog.askdirectory(
            title="Choose folder for experiment data",
            initialdir=self.exp_save_dir)
        if folder:
            self.exp_save_dir = folder
            self.exp_save_dir_var.set(folder)

    def _collect_params(self):
        """Validate the Experiment tab. Returns the params dict, or None after
        telling the user what is wrong."""
        errors = []

        def num(var, label, low=0.0, high=None):
            try:
                value = float(var.get())
            except (tk.TclError, ValueError):
                errors.append(f"- {label}: not a number")
                return 0.0
            if value < low or (high is not None and value > high):
                limit = (f"between {low:g} and {high:g}" if high is not None
                         else f"at least {low:g}")
                errors.append(f"- {label}: must be {limit} (got {value:g})")
            return value

        kp_max, ki_max, kd_max = MOTOR_GAIN_LIMITS
        rate_lo, rate_hi = SAMPLE_RATE_LIMITS
        params = {
            "name": self.exp_name_var.get().strip() or "fred_experiment",
            "heater_mode": self.exp_heater_mode_var.get(),
            "target_temperature": num(self.exp_target_temp_var, "Target temp"),
            "temp_kp": num(self.exp_temp_kp_var, "Temp Kp"),
            "temp_ki": num(self.exp_temp_ki_var, "Temp Ki"),
            "temp_kd": num(self.exp_temp_kd_var, "Temp Kd"),
            "heater_pwm": num(self.exp_heater_pwm_var, "Heater PWM", 0, 100),
            "spooler_mode": self.exp_spooler_mode_var.get(),
            "motor_setpoint": num(self.exp_motor_setpoint_var,
                                  "Spooler setpoint"),
            "motor_kp": num(self.exp_motor_kp_var, "Motor Kp", 0, kp_max),
            "motor_ki": num(self.exp_motor_ki_var, "Motor Ki", 0, ki_max),
            "motor_kd": num(self.exp_motor_kd_var, "Motor Kd", 0, kd_max),
            "dc_motor_pwm": num(self.exp_dc_pwm_var, "DC Motor PWM", 0, 100),
            "extrusion_speed": num(self.exp_extrusion_var, "Extrusion speed"),
            "fan_duty": num(self.exp_fan_var, "Fan duty", 0, 100),
            "target_diameter": num(self.exp_target_diameter_var,
                                   "Target diameter"),
            "heating_delay": num(self.exp_heating_delay_var, "Heating time"),
            "heat_extrude_time": num(self.exp_heat_extrude_time_var,
                                     "Heating + extrusion time"),
            "heat_extrude_speed": num(self.exp_heat_extrude_speed_var,
                                      "Extrusion rate during heating"),
            "data_delay": num(self.exp_data_delay_var, "Settle time"),
            "data_taking_time": num(self.exp_data_time_var, "Data-taking time"),
            "post_spool_time": num(self.exp_post_spool_var,
                                   "Extra spooling time"),
            "sample_rate_hz": num(self.exp_rate_var, "FrED sample rate",
                                  rate_lo, rate_hi),
        }
        if not errors and params["data_taking_time"] <= 0:
            errors.append("- Data-taking time: must be more than 0")
        if errors:
            messagebox.showwarning("Experiment settings",
                                   "Please fix these values:\n\n"
                                   + "\n".join(errors))
            return None
        return params

    def _new_run(self, params):
        return {"name": params.get("name", "fred_experiment"),
                "params": params, "sent_t": time.perf_counter(),
                "t0_pi": None, "t0_laptop": None, "t0_arrival": None,
                "boot": None, "steady_t": None, "start_now_t": None}

    def send_experiment(self):
        if not self.link.is_open:
            messagebox.showwarning(
                "Not connected", "Connect to FrED on the 'Measure & Connect' "
                "tab first.")
            return
        params = self._collect_params()
        if params is None:
            return
        if not self.calib.is_calibrated:
            if not messagebox.askyesno(
                "Not calibrated",
                "The camera is not calibrated, so diameter will be recorded in "
                "pixels, not mm. Send the experiment anyway?"):
                return
        if self.link.send({"type": "experiment", "params": params}):
            self.run = self._new_run(params)
            total = (params["heating_delay"] + params["heat_extrude_time"]
                     + params["data_delay"] + params["data_taking_time"])
            spool = params["post_spool_time"]
            spool_note = f" + {spool:.0f}s spooling" if spool else ""
            self.exp_status_var.set(
                f"Experiment sent. Heating... (total ~{total:.0f}s{spool_note})")
            self.status_var.set("Experiment sent to FrED.")
        else:
            self._mark_link_lost()
            messagebox.showerror("Send failed",
                                 "Could not send the experiment (link lost).")

    def start_recording_now(self):
        """Fiber dropped: FrED starts recording (and all systems) right now.

        During the warm-up of a run this skips the rest of it; with no run
        active it starts the Experiment tab's run directly in RECORDING.
        """
        if not self.link.is_open:
            messagebox.showwarning(
                "Not connected", "Connect to FrED on the 'Measure & Connect' "
                "tab first.")
            return
        if self.pi_phase in ("recording", "spooling"):
            self.status_var.set("FrED is already past the warm-up - "
                                "START RECORDING NOW ignored.")
            return
        params = self._collect_params()
        if params is None:
            return
        pressed = time.perf_counter()
        if not self.link.send({"type": "start_now", "params": params}):
            self._mark_link_lost()
            messagebox.showerror("Send failed",
                                 "Could not reach FrED (link lost).")
            return
        in_warm_up = (self.run is not None
                      and self.pi_phase in self.WARM_UP_PHASES)
        if not in_warm_up:
            self.run = self._new_run(params)
        self.run["start_now_t"] = pressed
        note = ("" if self.calib.is_calibrated
                else "  (camera NOT calibrated - diameter in pixels)")
        self.status_var.set("START RECORDING NOW sent to FrED." + note)

    def mark_steady_state(self):
        """Operator: the flow looks steady from now on. Only marks the data."""
        if self.run is None:
            messagebox.showinfo(
                "Steady state", "Send or start an experiment first - the mark "
                "is saved with that run's data.")
            return
        moved = self.run.get("steady_t") is not None
        self.run["steady_t"] = time.perf_counter()
        when = ""
        if self.run.get("t0_laptop") is not None:
            when = (f" at t = {self.run['steady_t'] - self.run['t0_laptop']:.2f}"
                    " s")
        self.status_var.set(("Steady-state mark moved" if moved else
                             "Steady state marked") + when +
                            " (press again to move it).")
        self._update_run_status()

    def abort_experiment(self):
        if not self.link.is_open:
            messagebox.showwarning("Not connected", "Not connected to FrED.")
            return
        if messagebox.askyesno("Abort", "Abort the running experiment on FrED?\n"
                               "Every system (heater, stepper, spooler, fan) "
                               "stops."):
            self.link.send({"type": "abort"})
            self.exp_status_var.set("Abort sent.")

    def _update_run_status(self):
        """The one-line FrED status under the graph."""
        if not self.link.is_open:
            self.run_status_var.set("FrED: not connected")
            return
        parts = [f"FrED: {self._phase_text or self.pi_phase}"]
        if self._phase_remaining is not None:
            remaining, t_rx = self._phase_remaining
            left = max(0.0, remaining - (time.perf_counter() - t_rx))
            parts[0] += f" ({left:.0f}s left)"
        if self.run is not None and self.run.get("steady_t") is not None:
            if self.run.get("t0_laptop") is not None:
                rel = self.run["steady_t"] - self.run["t0_laptop"]
                parts.append(f"steady state marked at t = {rel:.1f} s")
            else:
                parts.append("steady state marked")
        est = self.sync.estimate()
        if est is not None:
            parts.append(f"sync ±{est['err'] * 1000:.1f} ms")
        self.run_status_var.set("   |   ".join(parts))

    RETRIEVE_TIMEOUT_S = 60   # give up if FrED sends nothing for this long

    def retrieve_data(self):
        if not self.link.is_open:
            messagebox.showwarning("Not connected", "Connect to FrED first.")
            return
        if self._retrieve is not None:
            return                       # a retrieval is already running
        self._send_ping()                # one more sync sample, just in case
        if not self.link.send({"type": "get_data"}):
            self._mark_link_lost()
            messagebox.showerror("Request failed",
                                 "Could not ask FrED for the data (link lost).")
            return
        self.status_var.set("Requested experiment data from FrED...")
        self._open_retrieve_dialog()

    # ------------------------------------------------------------------ #
    # Modal progress dialog: blocks the app while the data is received,
    # processed and saved, and shows a loading bar for each stage.
    # ------------------------------------------------------------------ #
    def _open_retrieve_dialog(self):
        dlg = tk.Toplevel(self.root)
        dlg.title("Retrieving data from FrED")
        dlg.resizable(False, False)
        dlg.transient(self.root)
        frm = ttk.Frame(dlg, padding=(18, 14))
        frm.pack(fill=tk.BOTH, expand=True)
        label_var = tk.StringVar(
            value="Waiting for FrED to send the recorded data...")
        ttk.Label(frm, textvariable=label_var, wraplength=340).pack(anchor="w")
        bar = ttk.Progressbar(frm, mode="indeterminate", length=340)
        bar.pack(fill=tk.X, pady=(10, 8))
        bar.start(12)
        cancel_btn = ttk.Button(frm, text="Cancel",
                                command=self._cancel_retrieve)
        cancel_btn.pack()
        dlg.protocol("WM_DELETE_WINDOW", self._cancel_retrieve)
        # Center the dialog over the app, then make it modal (input to the
        # rest of the application is blocked until it closes).
        dlg.update_idletasks()
        x = self.root.winfo_rootx() + (self.root.winfo_width()
                                       - dlg.winfo_width()) // 2
        y = self.root.winfo_rooty() + (self.root.winfo_height()
                                       - dlg.winfo_height()) // 2
        dlg.geometry(f"+{max(x, 0)}+{max(y, 0)}")
        dlg.grab_set()
        self._retrieve = {
            "dlg": dlg, "bar": bar, "label": label_var, "cancel": cancel_btn,
            "phase": "waiting",       # waiting -> saving -> done
            "frac": 0.0, "text": "", "result": None, "t0": time.time(),
        }
        self.root.after(100, self._poll_retrieve)

    def _poll_retrieve(self):
        """Tk-thread heartbeat: updates the bar, enforces the timeout, and
        closes the dialog when the worker thread reports completion."""
        st = self._retrieve
        if st is None:
            return
        if st["phase"] == "waiting":
            elapsed = time.time() - st["t0"]
            if not self.link.is_open:
                self._mark_link_lost()
                self._finish_retrieve(
                    ("error", "Link to FrED lost while waiting for the data.\n"
                              "Reconnect on the 'Measure & Connect' tab, then "
                              "try Retrieve Data again."))
                return
            if elapsed > self.RETRIEVE_TIMEOUT_S:
                # The stream may be poisoned by a half-sent message - reset
                # the connection so the next attempt starts clean.
                if self._reset_link():
                    extra = ("\n\nThe WiFi link was reset - click Retrieve "
                             "Data to try again.")
                else:
                    extra = ("\n\nThe link could not be re-opened - reconnect "
                             "on the 'Measure & Connect' tab, then try again.")
                self._finish_retrieve(
                    ("error", f"FrED did not send the data within "
                              f"{self.RETRIEVE_TIMEOUT_S} seconds.{extra}"))
                return
            st["label"].set("Waiting for FrED to send the recorded data... "
                            f"({elapsed:.0f} s)")
        elif st["phase"] == "saving":
            st["bar"]["value"] = st["frac"] * 100
            st["label"].set(st["text"])
        elif st["phase"] == "done":
            self._finish_retrieve(st["result"])
            return
        self.root.after(100, self._poll_retrieve)

    def _cancel_retrieve(self):
        st = self._retrieve
        if st is None or st["phase"] != "waiting":
            return   # saving to disk cannot be cancelled - it finishes
        self._finish_retrieve(("cancelled", None))

    def _finish_retrieve(self, result):
        st, self._retrieve = self._retrieve, None
        if st is not None:
            try:
                st["dlg"].grab_release()
                st["dlg"].destroy()
            except tk.TclError:
                pass
        kind, payload = result
        if kind == "ok":
            saved, note = payload
            if self.run is not None:
                self.run["retrieved"] = True
            self.exp_status_var.set("Experiment data saved.")
            self.status_var.set(f"Saved {saved.splitlines()[0]}")
            messagebox.showinfo("Experiment data saved",
                                f"Saved FrED experiment data to:\n{saved}{note}")
        elif kind == "error":
            self.exp_status_var.set("Data retrieval failed.")
            messagebox.showerror("Retrieve data", payload)
        else:   # cancelled
            self.status_var.set("Data retrieval cancelled.")

    # ------------------------------------------------------------------ #
    # Messages coming back from FrED (run on the Tk main thread)
    # ------------------------------------------------------------------ #
    def _handle_pi_message(self, msg):
        mtype = msg.get("type")
        if mtype == "sync_reply":
            self.sync.add_reply(msg)
        elif mtype == "status":
            phase = msg.get("phase", "")
            text = msg.get("message", phase)
            remaining = msg.get("remaining", 0) or 0
            self.pi_phase = phase
            self._phase_text = text
            self._phase_remaining = ((float(remaining), time.perf_counter())
                                     if remaining else None)
            if phase == "recording" and msg.get("t0") is not None:
                if self.run is None:
                    self.run = self._new_run({})
                arrival = time.perf_counter()
                self.run["t0_pi"] = float(msg["t0"])
                self.run["t0_arrival"] = arrival
                self.run["boot"] = self.sync.last_boot
                self.run["t0_laptop"] = self._pi_to_laptop(
                    self.run["t0_pi"], self.run["boot"], arrival)[0]
            extra = f" ({remaining:.0f}s left)" if remaining else ""
            self.exp_status_var.set(f"Experiment: {text}{extra}")
            if msg.get("data_ready") or phase == "complete":
                self.exp_status_var.set(
                    f"Experiment: {text} - click 'Retrieve Data'.")
            self._update_run_status()
        elif mtype == "event" and msg.get("event") == "no_data":
            text = msg.get("message", "No experiment data available yet. "
                                      "Run an experiment first.")
            if self._retrieve is not None and self._retrieve["phase"] == "waiting":
                self._finish_retrieve(("error", text))
            else:
                messagebox.showinfo("No data", text)
        elif mtype == "data":
            self._receive_experiment_data(msg)

    def _receive_experiment_data(self, msg):
        st = self._retrieve
        if st is None or st["phase"] != "waiting":
            # Data arriving with no retrieval pending (e.g. after Cancel).
            self.status_var.set("Data from FrED arrived after cancel - "
                                "ignored. Click Retrieve Data to ask again.")
            return
        try:
            csv_text = base64.b64decode(msg.get("b64", "")).decode("utf-8")
        except Exception as exc:
            self._finish_retrieve(("error", f"Could not decode the data: {exc}"))
            return
        base = msg.get("name") or self.exp_name_var.get().strip() or "fred_experiment"
        for ch in '<>:"/\\|?*':
            base = base.replace(ch, "_")
        merge = self._prepare_merge(msg.get("meta") or {})
        # Switch the dialog to a determinate loading bar and process/save the
        # files on a worker thread so the bar keeps moving (the Excel build
        # takes several seconds for a long run).
        st["phase"] = "saving"
        st["cancel"].configure(state="disabled")
        st["bar"].stop()
        st["bar"].configure(mode="determinate", maximum=100, value=0)
        st["text"] = "Data received - merging with the camera data..."
        st["label"].set(st["text"])
        threading.Thread(target=self._retrieve_worker,
                         args=(st, base, csv_text, merge), daemon=True).start()

    def _prepare_merge(self, meta):
        """Tk thread: everything the merge needs (sync, window, camera data)."""
        run = self.run
        t0_pi, t_end_pi = meta.get("t0"), meta.get("t_end")
        boot = meta.get("boot")
        t0_laptop, method = None, "no timing information from FrED"
        if t0_pi is not None:
            fallback = None
            if run is not None and run.get("t0_pi") is not None \
                    and abs(run["t0_pi"] - float(t0_pi)) < 1e-6:
                fallback = run.get("t0_arrival")
            t0_laptop, method = self._pi_to_laptop(float(t0_pi), boot, fallback)
            if t0_laptop is None:
                method = ("no clock sync and no recording-start message - "
                          "camera data could not be placed on FrED's timeline")
        cam_rows = []
        if t0_laptop is not None and t_end_pi is not None:
            duration = float(t_end_pi) - float(t0_pi)
            cam_rows = self.log.since(t0_laptop - 2.0,
                                      t0_laptop + duration + 2.0)
        steady_rel = None
        if run is not None and run.get("steady_t") is not None \
                and t0_laptop is not None:
            steady_rel = run["steady_t"] - t0_laptop
        est = self.sync.estimate(boot) if boot else None
        start_now = (run is not None and run.get("start_now_t") is not None)
        return {"meta": meta, "t0_laptop": t0_laptop, "method": method,
                "cam_rows": cam_rows, "steady_rel": steady_rel, "sync": est,
                "params": dict(run["params"]) if run else {},
                "start_now": start_now}

    def _retrieve_worker(self, st, base, csv_text, merge):
        """Background thread: merge FrED's table with the camera data and
        write the CSVs + formatted Excel, reporting progress. No Tk calls in
        here - it only writes plain fields in ``st`` that the
        _poll_retrieve() heartbeat displays."""
        def progress(frac, text):
            st["frac"] = frac
            st["text"] = text
        try:
            folder = self.exp_save_dir
            os.makedirs(folder, exist_ok=True)
            csv_path = os.path.join(folder, base + ".csv")
            cam_path = os.path.join(folder, base + "_camera.csv")
            xlsx_path = os.path.join(folder, base + ".xlsx")
            progress(0.02, "Merging FrED's data with the camera data...")
            fred_header, fred_rows = parse_fred_csv(csv_text)
            notes = []
            if "Temp new reading" not in fred_header:
                # FrED runs pre-v7 software: its table already has its own
                # diameter columns; save it as it came.
                main_header, main_rows = fred_header, fred_rows
                cam_header, cam_rows, stats = [], [], None
                notes.append("FrED is running older software (no camera "
                             "merge) - run 'git pull' on the Pi.")
                with open(csv_path, "w", newline="", encoding="utf-8") as fh:
                    fh.write(csv_text)
            else:
                (main_header, main_rows, cam_header, cam_rows,
                 stats) = merge_run(fred_header, fred_rows, merge["cam_rows"],
                                    merge["t0_laptop"], merge["steady_rel"])
                progress(0.04, "Saving CSV files...")
                with open(csv_path, "w", newline="", encoding="utf-8") as fh:
                    fh.write(format_semicolon_csv(main_header, main_rows))
                if cam_rows:
                    with open(cam_path, "w", newline="", encoding="utf-8") as fh:
                        fh.write(format_semicolon_csv(cam_header, cam_rows))
                else:
                    notes.append("No camera data covered the recording, so the "
                                 "diameter columns are empty (was the camera "
                                 "running and this app open during the run?).")
            info = self._run_info(base, merge, stats)
            xlsx_ok, xlsx_msg = write_run_xlsx(
                xlsx_path, main_header, main_rows, cam_header, cam_rows, info,
                progress=progress)
            progress(1.0, "Done.")
            saved = csv_path
            if cam_rows:
                saved += "\n" + cam_path
            if xlsx_ok:
                saved += "\n" + xlsx_path
            else:
                notes.append(f"Excel not written ({xlsx_msg}).")
            if stats:
                notes.insert(0, (
                    f"Diameter: {stats['frames']} camera frames "
                    f"({stats['fps']:.1f} fps) merged onto {stats['rows']} "
                    f"FrED rows; time alignment by {merge['method']}."))
            note = ("\n\n" + "\n".join(notes)) if notes else ""
            st["result"] = ("ok", (saved, note))
        except Exception as exc:
            st["result"] = ("error", f"Could not save the data:\n{exc}")
        st["phase"] = "done"

    @staticmethod
    def _run_info(base, merge, stats):
        """Rows of the 'Run info' sheet."""
        meta = merge["meta"]
        est = merge["sync"]
        info = [
            ("Experiment", base),
            ("Saved", _dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
            ("Software", f"FrED Fiber Measure {APP_VERSION}"),
            ("t = 0", "FrED recording start ("
                      + ("START RECORDING NOW button" if merge["start_now"]
                         else "experiment timer") + ")"),
            ("Time alignment", merge["method"]),
        ]
        if est is not None:
            info += [
                ("Clock sync pings", est["n"]),
                ("Clock offset Pi - laptop (s)", round(est["a"], 6)),
                ("Clock drift (ppm)", round(est["b"] * 1e6, 3)),
                ("Estimated alignment error (ms)", round(est["err"] * 1000, 2)),
            ]
        if meta:
            info += [
                ("FrED rows", meta.get("rows")),
                ("FrED rate achieved (Hz)", round(meta.get("rate_hz") or 0, 2)),
            ]
        if stats:
            info += [
                ("Recording duration (s)", round(stats["duration"], 3)),
                ("Camera frames in recording", stats["frames"]),
                ("Camera rate (fps)", round(stats["fps"], 2)),
                ("Frames with fiber detected", stats["detected"]),
                ("FrED rows with a diameter", stats["rows_with_diameter"]),
                ("FrED rows with a NEW camera frame", stats["new_frame_rows"]),
                ("Diameter unit", stats["unit"]),
            ]
        steady = merge["steady_rel"]
        info += [
            ("Steady state marked at (s)",
             round(steady, 3) if steady is not None else "not marked"),
            ("Diameter columns",
             "latest camera frame at or before each FrED row (never "
             "interpolated); 'Diameter new frame' 1 = first row with that "
             "frame, 0 = repeated to fill the row"),
            ("Diameter filter",
             f"centred median of {FILTER_MEDIAN} + mean of {FILTER_MEAN} "
             "detected frames (no time lag); 'raw' = as measured"),
            ("Temp / Spooler new reading",
             "1 = fresh sensor reading in that row, 0 = repeated value"),
            ("Camera latency compensation (s)", CAMERA_LATENCY_S),
        ]
        for key, value in sorted((merge.get("params") or {}).items()):
            info.append((f"Param: {key}", value))
        return info

    # ------------------------------------------------------------------ #
    # Shutdown
    # ------------------------------------------------------------------ #
    def on_close(self):
        if self.recording and self.records:
            if not messagebox.askyesno(
                "Quit",
                "A recording is still running and has unsaved data.\n"
                "Quit anyway? (unsaved data will be lost)",
            ):
                return
        if (self.run is not None and self.run.get("t0_pi") is not None
                and not self.run.get("retrieved")):
            if not messagebox.askyesno(
                "Quit",
                "The camera data of the last FrED experiment lives in this "
                "app until you click Retrieve Data.\n"
                "Quit anyway? (if not retrieved yet, its diameter data will "
                "be lost)",
            ):
                return
        self.worker.stop()
        self.link.close()
        self.root.destroy()


def main():
    root = tk.Tk()
    FiberApp(root)
    root.mainloop()


if __name__ == "__main__":
    main()
